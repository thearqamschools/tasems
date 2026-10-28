"""Shared parts of the Excel imports: the rows that can still be saved when others have
problems, and the Excel list of problem rows to correct and import again."""

import io

from django.http import HttpResponse
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def problem_row_numbers(errors):
    """Rows with at least one problem. errors are (row number or None, column, message)."""
    return {number for number, _col, _msg in errors if number is not None}


def file_errors(errors):
    """Problems with the whole file (missing columns, too many rows); these block any import."""
    return [e for e in errors if e[0] is None]


def problems_response(header, raw_rows, errors, filename):
    """An Excel file of the rows with problems, in the original columns plus a "Problems" column,
    so they can be corrected and imported again on their own."""
    wb = Workbook()
    ws = wb.active
    ws.title = _("Problem rows")
    ws.append([str(h or "") for h in header] + [_("Problems"), _("Row in original file")])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E3EFEB")
    messages = {}
    for number, column, message in errors:
        if number is not None:
            messages.setdefault(number, []).append(f"{column}: {message}" if column else message)
    width = len(header)
    for number in sorted(messages):
        values = list(raw_rows.get(number, ()))[:width]
        values += [None] * (width - len(values))
        ws.append(values + ["; ".join(messages[number]), number])
    ws.column_dimensions[ws.cell(1, width + 1).column_letter].width = 60
    ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type=XLSX)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
