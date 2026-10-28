"""Excel import of students for one branch.

Step 1 (check): read the file, report every problem by row, and list the sections
that will be created. Nothing is saved.
Step 2 (import): when the check finds no problems, or, if the user chooses to, only the
rows without problems (the rest are skipped and can be downloaded to fix). Everything is
saved in one transaction, so either all chosen rows go in or none do.

Two kinds of file are accepted: the TASEMS Excel template, and the "Student Data" PDF report
printed by the school's previous software (one class/section per page, with Comp. Code,
Student Name, Father Name, Class, Section, Mon. Fee, Mobile, Other Number, CNIC/Occupation).
"""

import datetime as dt
import io
import re
from dataclasses import dataclass, field

from django.db import transaction
from django.utils.translation import gettext as _
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from apps.academics.models import ClassLevel, Section
from apps.core.importing import file_errors, problem_row_numbers
from apps.core.phone import normalize_cnic, normalize_mobile

from .models import Student, StudentGuardian
from .services import save_student

MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024

# key, header, required
COLUMNS = [
    ("admission_no", "Admission No", True),
    ("first_name", "First Name", True),
    ("last_name", "Last Name", False),
    ("full_name_ur", "Name (Urdu)", False),
    ("gender", "Gender", True),
    ("date_of_birth", "Date of Birth", False),
    ("admission_date", "Admission Date", False),
    ("b_form_no", "B-Form No", False),
    ("class", "Class", True),
    ("section", "Section", True),
    ("is_hifz", "Hifz", False),
    ("guardian_name", "Parent Name", True),
    ("guardian_name_ur", "Parent Name (Urdu)", False),
    ("guardian_phone", "Parent Mobile", True),
    ("relationship", "Relationship", False),
    ("father_name", "Father Name", False),
    ("father_cnic", "Father CNIC", False),
    ("alt_phone", "Other Mobile", False),
    ("occupation", "Father Occupation", False),
    ("monthly_fee", "Monthly Fee", False),
]
HEADER_TO_KEY = {re.sub(r"[^a-z]", "", h.lower()): k for k, h, _r in COLUMNS}

GENDERS = {
    "male": "male",
    "m": "male",
    "boy": "male",
    "لڑکا": "male",
    "female": "female",
    "f": "female",
    "girl": "female",
    "لڑکی": "female",
}
YES = {"yes", "y", "1", "true", "ہاں", "جی"}
RELATIONSHIPS = {
    "father": "father",
    "والد": "father",
    "mother": "mother",
    "والدہ": "mother",
    "guardian": "guardian",
    "سرپرست": "guardian",
    "other": "other",
}
CLASS_ALIASES = {
    "one": "class-1",
    "two": "class-2",
    "three": "class-3",
    "four": "class-4",
    "five": "class-5",
    "six": "class-6",
    "seven": "class-7",
    "eight": "class-8",
    "nine": "class-9",
    "ten": "class-10",
    "matric": "class-10",
    "pg": "play-group",
    "playgroup": "play-group",
    "nur": "nursery",
    "nursery": "nursery",
    "prep": "prep",
    "pre9": "pre-9",
    "prenine": "pre-9",
    "hifz": "hifz",
}


def _key(value):
    return re.sub(r"[^a-z0-9]", "", str(value).lower())


def class_lookup():
    table = {}
    for level in ClassLevel.objects.filter(is_active=True):
        table[_key(level.code)] = level
        table[_key(level.name)] = level
        number = re.fullmatch(r"class-(\d+)", level.code)
        if number:
            table[number.group(1)] = level
    for alias, code in CLASS_ALIASES.items():
        level = next((lvl for lvl in table.values() if lvl.code == code), None)
        if level:
            table[alias] = level
    return table


def parse_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%b-%Y", "%d %b %Y"):
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


@dataclass
class Row:
    number: int
    data: dict
    level: ClassLevel | None = None
    defaulted_admission: bool = False


@dataclass
class CheckResult:
    rows: list = field(default_factory=list)
    errors: list = field(default_factory=list)  # (row number or None, column header or "", message)
    new_sections: list = field(default_factory=list)  # (level, name, gender_group)
    header: tuple = ()
    raw: dict = field(default_factory=dict)  # row number -> values as in the file

    @property
    def ok(self):
        return not self.errors and bool(self.rows)

    @property
    def problem_rows(self):
        return problem_row_numbers(self.errors)

    @property
    def good_rows(self):
        bad = self.problem_rows
        return [r for r in self.rows if r.number not in bad]

    @property
    def can_skip_problems(self):
        """Some rows have problems but the rest can be imported on their own."""
        return bool(self.problem_rows) and not file_errors(self.errors) and bool(self.good_rows)

    @property
    def good_new_sections(self):
        used = {(r.level.id, r.data["section"].lower()) for r in self.good_rows}
        return [s for s in self.new_sections if (s[0].id, s[1].lower()) in used]

    @property
    def defaulted_admission_dates(self):
        return sum(r.defaulted_admission for r in self.good_rows)


def template_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.append([h + (" *" if required else "") for _k, h, required in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E3EFEB")
    for col, width in zip("ABCDEFGHIJKLMNO", [14, 16, 16, 20, 10, 14, 15, 16, 12, 10, 8, 20, 22, 16, 13], strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    guide = wb.create_sheet("How to fill")
    lines = [
        ("Columns marked * are required. Keep the header row exactly as it is.",),
        ("Gender", "Male or Female (also accepted: Boy, Girl, M, F)"),
        ("Dates", "Day first: 15/04/2026 or 15-04-2026, or an Excel date"),
        ("Class", "Play Group, Nursery, Prep, Class 1 … Class 8, Pre-9, Class 9, Class 10, Hifz (1–10 also work)"),
        ("Section", "The section name, e.g. A, B, Rose. Missing sections are created for the current session"),
        ("Hifz", "Yes for Hifz-e-Quran students, otherwise leave empty"),
        ("Parent Mobile", "Pakistani mobile, e.g. 0300-1234567"),
        ("Relationship", "Father, Mother, Guardian or Other (empty = Guardian)"),
        ("From Class 6", "Boys and girls must be in separate sections"),
        ("Admission Date", "If empty, the first day of the session is used (correct it later)"),
        ("Monthly Fee", "The child's agreed monthly fee, if different families pay different amounts"),
        ("Old software", "The previous software's 'Student Data' PDF can also be uploaded as it is"),
    ]
    for line in lines:
        guide.append(list(line))
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 90
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


OLD_REPORT_COLUMNS = [
    "Admission No",
    "First Name",
    "Gender",
    "Class",
    "Section",
    "Parent Name",
    "Parent Mobile",
    "Relationship",
    "Father Name",
    "Father CNIC",
    "Other Mobile",
    "Father Occupation",
    "Monthly Fee",
]


def _old_report_rows(file_bytes):
    """Rows of the previous software's "Student Data" PDF, in TASEMS column names."""
    import pdfplumber

    rows, gender = [], ""
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            campus = re.search(r"Campus\s*\((\w+)\)", text)
            if campus:  # a class continued on the next page keeps the campus of its first page
                gender = {"boys": "Male", "girls": "Female"}.get(campus.group(1).lower(), "")
            for table in page.extract_tables():
                head = [re.sub(r"[^a-z]", "", str(c or "").lower()) for c in table[0]]
                if "compcode" not in head:
                    continue
                col = {name: head.index(name) for name in head}
                for raw in table[1:]:

                    def cell(name, _raw=raw, _col=col):
                        i = _col.get(name)
                        return str(_raw[i] or "").strip() if i is not None and i < len(_raw) else ""

                    code = cell("compcode")
                    if not code:
                        continue
                    numbers = [
                        n for n in re.split(r"[\n,/]+", cell("mobile") + "\n" + cell("othernumber")) if n.strip()
                    ]
                    mobile = next((n for n in numbers if normalize_mobile(n)), numbers[0] if numbers else "")
                    others = [n.strip() for n in numbers if n != mobile]
                    cnic_text = cell("cnicoccupation")
                    found = re.search(r"\d{5}-?\d{7}-?\d", cnic_text)
                    occupation = re.sub(
                        r"\s+", " ", cnic_text.replace(found.group(0), "") if found else cnic_text
                    ).strip()
                    father = re.sub(r"\s+", " ", cell("fathername"))
                    rows.append(
                        [
                            code,
                            re.sub(r"\s+", " ", cell("studentname")),
                            gender,
                            re.sub(r"\s+", " ", cell("class")),
                            re.sub(r"\s+", " ", cell("section")),
                            father,
                            mobile,
                            "Father",
                            father,
                            found.group(0) if found else "",
                            (others[0] if others else "")[:20],
                            occupation[:100],
                            cell("monfee"),
                        ]
                    )
    return OLD_REPORT_COLUMNS, rows


def _table(file_bytes):
    """(header, rows) from an Excel template or the previous software's PDF report."""
    if file_bytes[:4] == b"%PDF":
        header, rows = _old_report_rows(file_bytes)
        return header, iter(rows)
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    rows = wb.worksheets[0].iter_rows(values_only=True)
    return next(rows, None) or (), rows


def check(file_bytes, *, branch, session):
    result = CheckResult()
    if session is None:
        result.errors.append((None, "", _("No current academic session is set. Ask head office to set it first.")))
        return result
    try:
        header, rows = _table(file_bytes)
        result.header = tuple(header)
    except Exception:  # noqa: BLE001 - any unreadable file gets the same message
        result.errors.append(
            (
                None,
                "",
                _("This file could not be opened. Use the Excel template (.xlsx) or the old software's PDF report."),
            )
        )
        return result
    positions = {}
    for index, title in enumerate(header):
        key = HEADER_TO_KEY.get(re.sub(r"[^a-z]", "", str(title or "").lower()))
        if key:
            positions[key] = index
    missing = [h for k, h, required in COLUMNS if required and k not in positions]
    if missing:
        result.errors.append(
            (
                None,
                "",
                _("These required columns are missing: %(cols)s. Use the template.") % {"cols": ", ".join(missing)},
            )
        )
        return result

    classes = class_lookup()
    headers = {k: h for k, h, _r in COLUMNS}
    seen_admission = {}
    existing = {a.lower() for a in Student.objects.filter(branch=branch).values_list("admission_no", flat=True)}
    sections = {(s.class_level_id, s.name.lower()): s for s in Section.objects.filter(branch=branch, session=session)}
    planned_genders = {}

    for number, values in enumerate(rows, start=2):
        if values is None or all(v in (None, "") for v in values):
            continue
        if len(result.rows) >= MAX_ROWS:
            result.errors.append(
                (None, "", _("Too many rows. Import at most %(n)s students at a time.") % {"n": MAX_ROWS})
            )
            break
        result.raw[number] = tuple(values)
        raw = {k: (values[i] if i < len(values) else None) for k, i in positions.items()}
        data, errors, defaulted = {}, [], False

        def err(column, message, _number=number, _errors=errors):
            _errors.append((_number, headers[column], message))

        for key, _h, required in COLUMNS:
            if key in ("date_of_birth", "admission_date"):
                continue
            data[key] = _text(raw.get(key))
            if required and not data[key]:
                err(key, _("Required"))

        adm = data["admission_no"]
        if adm:
            if adm.lower() in existing:
                err("admission_no", _("Already exists in this branch"))
            elif adm.lower() in seen_admission:
                err("admission_no", _("Repeated in this file (row %(r)s)") % {"r": seen_admission[adm.lower()]})
            seen_admission.setdefault(adm.lower(), number)

        if data["gender"]:
            gender = GENDERS.get(data["gender"].lower())
            if gender is None:
                err("gender", _("Use Male or Female"))
            data["gender"] = gender

        for key in ("date_of_birth", "admission_date"):
            try:
                data[key] = parse_date(raw.get(key))
            except ValueError:
                data[key] = None
                err(key, _("Not a date. Use day/month/year, e.g. 15/04/2026"))
            else:
                if key == "admission_date" and data[key] is None:
                    # Not in the old software's report: the session's first day, to correct later.
                    data[key] = session.starts_on
                    defaulted = True
        if data.get("date_of_birth") and data["date_of_birth"] > dt.date.today():
            err("date_of_birth", _("Date of birth is in the future"))

        level = classes.get(_key(data["class"])) if data["class"] else None
        if data["class"] and level is None:
            err("class", _("Unknown class. See the 'How to fill' sheet"))

        if data["guardian_phone"]:
            phone = normalize_mobile(data["guardian_phone"])
            if phone is None:
                err("guardian_phone", _("Not a Pakistani mobile number"))
            data["guardian_phone"] = phone

        relationship = data.get("relationship", "").lower()
        data["relationship"] = RELATIONSHIPS.get(relationship, "guardian" if not relationship else None)
        if data["relationship"] is None:
            err("relationship", _("Use Father, Mother, Guardian or Other"))
        data["is_hifz"] = data.get("is_hifz", "").lower() in YES
        fee = data.get("monthly_fee", "").replace(",", "")
        if fee:
            try:
                data["monthly_fee"] = int(float(fee))
                if data["monthly_fee"] < 0:
                    raise ValueError
            except ValueError:
                err("monthly_fee", _("Not an amount"))
                data["monthly_fee"] = None
        else:
            data["monthly_fee"] = None
        if data.get("father_cnic"):
            cnic = normalize_cnic(data["father_cnic"])
            if cnic is None:
                err("father_cnic", _("A CNIC has 13 digits, e.g. 33100-1234567-1"))
            data["father_cnic"] = cnic or ""
        if not data.get("father_name") and data["relationship"] == "father":
            data["father_name"] = data["guardian_name"]

        if level and data["section"] and data.get("gender"):
            key = (level.id, data["section"].lower())
            section = sections.get(key)
            if section is not None:
                group = section.gender_group
            elif level.separate_by_gender:
                group = "boys" if data["gender"] == "male" else "girls"
            else:
                group = "mixed"
            if group != "mixed" and {"boys": "male", "girls": "female"}[group] != data["gender"]:
                err(
                    "section",
                    _("Section %(s)s of %(c)s is for %(g)s only")
                    % {
                        "s": data["section"],
                        "c": level.name,
                        "g": Section.GenderGroup(group).label,
                    },
                )
            elif section is None:
                first = planned_genders.setdefault(key, (group, data["section"], level))
                if first[0] != group:
                    err(
                        "section",
                        _("Boys and girls cannot share section %(s)s of %(c)s")
                        % {"s": data["section"], "c": level.name},
                    )

        result.errors.extend(errors)
        result.rows.append(Row(number=number, data=data, level=level, defaulted_admission=defaulted))

    if not result.rows and not result.errors:
        result.errors.append((None, "", _("The file has no student rows.")))
    result.new_sections = [(level, name, group) for (group, name, level) in planned_genders.values()]
    return result


@transaction.atomic
def run_import(result, *, branch, session, user=None, skip_problems=False):
    """Saves every row, or with skip_problems only the rows without problems."""
    if not (result.ok or (skip_problems and result.can_skip_problems)):
        raise ValueError("Import called on a file with problems")
    sections = {(s.class_level_id, s.name.lower()): s for s in Section.objects.filter(branch=branch, session=session)}
    for level, name, group in result.good_new_sections:
        section = Section(branch=branch, session=session, class_level=level, name=name, gender_group=group)
        section.full_clean()
        section.save()
        sections[(level.id, name.lower())] = section
    rows = result.good_rows
    for row in rows:
        d = row.data
        student = Student(
            branch=branch,
            admission_no=d["admission_no"],
            first_name=d["first_name"],
            last_name=d["last_name"],
            full_name_ur=d["full_name_ur"],
            gender=d["gender"],
            date_of_birth=d["date_of_birth"],
            admission_date=d["admission_date"],
            b_form_no=d["b_form_no"],
            is_hifz=d["is_hifz"],
            father_name=d.get("father_name", ""),
            father_cnic=d.get("father_cnic", ""),
        )
        save_student(
            student,
            session=session,
            section=sections[(row.level.id, d["section"].lower())],
            roll_no=None,
            guardian_name=d["guardian_name"],
            guardian_name_ur=d["guardian_name_ur"],
            guardian_phone=d["guardian_phone"],
            relationship=d["relationship"] or StudentGuardian.Relationship.GUARDIAN,
        )
        link = student.guardian_links.filter(is_primary_contact=True).select_related("guardian").first()
        if link and (d.get("alt_phone") or d.get("occupation")):
            g = link.guardian
            g.alt_phone = g.alt_phone or d.get("alt_phone", "")
            g.occupation = g.occupation or d.get("occupation", "")
            g.save()
        if d.get("monthly_fee") is not None and user is not None:
            from apps.fees.services import set_agreed_fee

            set_agreed_fee(student, d["monthly_fee"], user, session=session)
    return len(rows)
