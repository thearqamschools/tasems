"""Admissions screens: inquiries, follow-ups, campus visits, conversion to a student, reports."""

import datetime as dt
import io

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import Http404, HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from django.utils.translation import gettext as _

from apps.academics.models import ClassLevel
from apps.core.branch import current_branch
from apps.organisation.models import AcademicSession
from apps.people.forms import StudentForm
from apps.people.services import save_student

from . import services as S
from .forms import FollowUpForm, InquiryForm
from .models import Inquiry, Source


def _branch(request, perm=S.VIEW_PERM):
    branch = current_branch(request)
    if branch is None or not request.user.has_branch_perm(perm, branch.id):
        raise PermissionDenied
    return branch


def _inquiry(branch, pk):
    inquiry = (
        Inquiry.objects.filter(pk=pk, branch=branch).select_related("class_level", "assigned_to", "student").first()
    )
    if inquiry is None:
        raise Http404
    return inquiry


def _date(value, default):
    try:
        return dt.date.fromisoformat(value) if value else default
    except ValueError:
        return default


def _filtered(request, branch):
    qs = Inquiry.objects.filter(branch=branch).select_related("class_level", "assigned_to")
    view = request.GET.get("view", "open")
    today = timezone.localdate()
    if view == "due":
        qs = qs.filter(status__in=Inquiry.OPEN, next_follow_up__lte=today)
    elif view == "open":
        qs = qs.filter(status__in=Inquiry.OPEN)
    status = request.GET.get("status")
    if status in Inquiry.Status.values:
        qs = qs.filter(status=status)
    level = request.GET.get("class")
    if level and level.isdigit():
        qs = qs.filter(class_level_id=level)
    source = request.GET.get("source")
    if source in Source.values:
        qs = qs.filter(source=source)
    if request.GET.get("mine"):
        qs = qs.filter(assigned_to=request.user)
    q = request.GET.get("q", "").strip()
    if q:
        digits = "".join(c for c in q if c.isdigit())
        cond = Q(child_name__icontains=q) | Q(parent_name__icontains=q)
        if len(digits) >= 4:
            cond |= Q(phone__contains=digits[-10:])
        qs = qs.filter(cond)
    order = "next_follow_up" if view in ("due", "open") else "-created_at"
    return qs.order_by(order, "-created_at"), view


@login_required
def index(request):
    branch = _branch(request)
    qs, view = _filtered(request, branch)
    today = timezone.localdate()
    open_qs = Inquiry.objects.filter(branch=branch, status__in=Inquiry.OPEN)
    start = today.replace(day=1)
    if request.GET.get("format") == "xlsx":
        return _xlsx(qs[:5000], f"Inquiries-{branch.code}-{today:%Y-%m-%d}.xlsx")
    return render(
        request,
        "admissions/index.html",
        {
            "page": Paginator(qs, 50).get_page(request.GET.get("page")),
            "view": view,
            "figures": {
                "due": open_qs.filter(next_follow_up__lte=today).count(),
                "open": open_qs.count(),
                "month": Inquiry.objects.filter(branch=branch, created_at__date__gte=start).count(),
                "admitted": Inquiry.objects.filter(branch=branch, status="admitted", closed_on__gte=start).count(),
            },
            "visits": open_qs.filter(visit_at__date=today).order_by("visit_at"),
            "classes": ClassLevel.objects.filter(is_active=True).order_by("sort_order"),
            "sources": Source.choices,
            "statuses": Inquiry.Status.choices,
            "query": request.GET.urlencode(),
            "can_add": request.user.has_branch_perm(S.ADD_PERM, branch.id),
            "today": today,
        },
    )


@login_required
def new(request):
    branch = _branch(request, S.ADD_PERM)
    form = InquiryForm(request.POST or None, branch=branch, initial={"assigned_to": request.user.id})
    if request.method == "POST" and form.is_valid():
        inquiry = S.create(form.save(commit=False), request.user)
        messages.success(request, _("Inquiry saved for %(n)s.") % {"n": inquiry.child_name})
        return redirect("admissions:detail", pk=inquiry.pk)
    return render(request, "admissions/form.html", {"form": form})


@login_required
def edit(request, pk):
    branch = _branch(request, S.CHANGE_PERM)
    inquiry = _inquiry(branch, pk)
    form = InquiryForm(request.POST or None, instance=inquiry, branch=branch)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, _("Saved."))
        return redirect("admissions:detail", pk=inquiry.pk)
    return render(request, "admissions/form.html", {"form": form, "inquiry": inquiry})


@login_required
def detail(request, pk):
    branch = _branch(request)
    inquiry = _inquiry(branch, pk)
    can_change = request.user.has_branch_perm(S.CHANGE_PERM, branch.id)
    initial = {"next_follow_up": timezone.localdate() + dt.timedelta(days=2), "kind": "call"}
    form = FollowUpForm(request.POST or None, initial=initial)
    if request.method == "POST":
        if not can_change:
            raise PermissionDenied
        if form.is_valid():
            d = form.cleaned_data
            visit_at = d.get("visit_at")
            if visit_at and timezone.is_naive(visit_at):
                visit_at = timezone.make_aware(visit_at)
            S.record(
                inquiry,
                request.user,
                d["kind"],
                d.get("note", ""),
                next_follow_up=d.get("next_follow_up"),
                status=d.get("status") or None,
                lost_reason=d.get("lost_reason", ""),
                visit_at=visit_at,
            )
            messages.success(request, _("Follow-up saved."))
            return redirect("admissions:detail", pk=inquiry.pk)
    earlier, at_school = S.matches(branch, inquiry.phone, exclude=inquiry)
    return render(
        request,
        "admissions/detail.html",
        {
            "inquiry": inquiry,
            "form": form,
            "timeline": inquiry.follow_ups.select_related("by"),
            "earlier": earlier,
            "at_school": at_school,
            "whatsapp": S.whatsapp_link(inquiry.phone, S.default_message(inquiry, branch)),
            "can_change": can_change,
            "can_admit": can_change and request.user.has_branch_perm("people.add_student", branch.id),
        },
    )


@login_required
def admit(request, pk):
    """Makes the student record from the inquiry (same form as Students → Add)."""
    branch = _branch(request, S.CHANGE_PERM)
    if not request.user.has_branch_perm("people.add_student", branch.id):
        raise PermissionDenied
    inquiry = _inquiry(branch, pk)
    if inquiry.student_id:
        return redirect("people:student_detail", pk=inquiry.student_id)
    session = AcademicSession.objects.filter(is_current=True).first()
    names = inquiry.child_name.split(" ", 1)
    initial = {
        "first_name": names[0],
        "last_name": names[1] if len(names) > 1 else "",
        "gender": inquiry.gender,
        "date_of_birth": inquiry.date_of_birth,
        "admission_date": timezone.localdate(),
        "is_hifz": inquiry.is_hifz,
        "guardian_name": inquiry.parent_name,
        "guardian_phone": inquiry.phone,
        "relationship": "father",
        "father_name": inquiry.parent_name,
    }
    form = StudentForm(request.POST or None, initial=initial, branch=branch, session=session)
    if session:
        form.fields["section"].queryset = form.fields["section"].queryset.filter(class_level=inquiry.class_level)
    if request.method == "POST" and form.is_valid():
        d = form.cleaned_data
        student = save_student(
            form.save(commit=False),
            session=session,
            section=d.get("section"),
            roll_no=d.get("roll_no"),
            guardian_name=d["guardian_name"],
            guardian_name_ur=d.get("guardian_name_ur", ""),
            guardian_phone=d["guardian_phone"],
            relationship=d["relationship"],
        )
        S.admit(inquiry, student, request.user)
        messages.success(
            request, _("%(n)s is now a student. Make the admission fee voucher next.") % {"n": student.full_name}
        )
        if request.user.has_branch_perm("fees.add_voucher", branch.id) and d.get("section"):
            return redirect("fees:issue", student_id=student.id)
        return redirect("people:student_detail", pk=student.pk)
    return render(request, "admissions/admit.html", {"form": form, "inquiry": inquiry, "session": session})


@login_required
def report(request):
    branch = _branch(request)
    today = timezone.localdate()
    start = _date(request.GET.get("from"), today.replace(day=1))
    end = _date(request.GET.get("to"), today)
    if end < start:
        start, end = end, start
    return render(request, "admissions/report.html", {"start": start, "end": end, **S.report(branch, start, end)})


def _xlsx(inquiries, filename):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Inquiries"
    headers = [
        "Date", "Child", "Gender", "Class", "Hifz", "Parent", "Mobile", "Area", "Source", "Status",
        "Reason (not joining)", "Next follow-up", "Follow-up by",
    ]  # fmt: skip
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E3EFEB")
    for i in inquiries:
        ws.append(
            [
                timezone.localtime(i.created_at).date(),
                i.child_name,
                i.get_gender_display(),
                i.class_level.name,
                "Yes" if i.is_hifz else "",
                i.parent_name,
                i.phone,
                i.area,
                i.get_source_display(),
                i.get_status_display(),
                i.get_lost_reason_display() if i.lost_reason else "",
                i.next_follow_up,
                (i.assigned_to.get_full_name() or i.assigned_to.username) if i.assigned_to else "",
            ]
        )
    for col, width in zip("ABCDEFGHIJKLM", [11, 22, 8, 12, 6, 22, 13, 20, 18, 16, 20, 13, 16], strict=True):
        ws.column_dimensions[col].width = width
    buffer = io.BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
