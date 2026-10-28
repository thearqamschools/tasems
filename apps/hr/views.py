"""Staff screens: register, attendance from the machines, corrections, leave, self-service."""

import datetime as dt
import io
import secrets
from pathlib import Path

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.db.models import Q
from django.http import FileResponse, Http404, HttpResponse
from django.shortcuts import redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.translation import gettext as _
from django.views.decorators.http import require_POST

from apps.core import exports
from apps.core.branch import current_branch
from apps.organisation.models import AcademicSession

from . import importer
from . import services as S
from .forms import CorrectionForm, EmployeeForm, LeaveForm
from .models import Employee, HRSettings, LeaveRequest, StaffDay


def _branch(request, perm=S.VIEW_PERM):
    branch = current_branch(request)
    if branch is None or not request.user.has_branch_perm(perm, branch.id):
        raise PermissionDenied
    return branch


def _employee(branch, employee_id):
    e = Employee.objects.filter(pk=employee_id, branch=branch).select_related("campus", "user").first()
    if e is None:
        raise Http404
    return e


def _date(value, default):
    try:
        return dt.date.fromisoformat(value) if value else default
    except ValueError:
        return default


def _month(request):
    today = timezone.localdate()
    try:
        y, m = (int(x) for x in request.GET.get("month", "").split("-"))
        return dt.date(y, m, 1)
    except (ValueError, TypeError):
        return today.replace(day=1)


def _session():
    return AcademicSession.objects.filter(is_current=True).first()


def _month_end(first):
    return (first.replace(day=28) + dt.timedelta(days=4)).replace(day=1) - dt.timedelta(days=1)


@login_required
def index(request):
    branch = _branch(request)
    today = timezone.localdate()
    qs = Employee.objects.filter(branch=branch).select_related("campus")
    status = request.GET.get("status", "active")
    if status in Employee.Status.values:
        qs = qs.filter(status=status)
    staff_type = request.GET.get("type")
    if staff_type in Employee.StaffType.values:
        qs = qs.filter(staff_type=staff_type)
    q = request.GET.get("q", "").strip()
    if q:
        qs = qs.filter(
            Q(full_name__icontains=q) | Q(code__iexact=q) | Q(phone__icontains=q) | Q(designation__icontains=q)
        )
    if exports.wants_xlsx(request):
        return exports.xlsx_response(
            _("Staff"),
            [_("Code"), _("Name"), _("Father / husband name"), _("Designation"), _("Type"), _("Mobile"), _("CNIC"),
             _("Campus"), _("Qualification"), _("Joining date"), _("Machine ID"), _("Status")],
            (
                [e.code, e.full_name, e.father_name or e.husband_name, e.designation, e.get_staff_type_display(),
                 e.phone, e.cnic, str(e.campus or ""), e.qualification, e.joined_on, e.machine_id,
                 e.get_status_display()]
                for e in qs.order_by("staff_type", "full_name")
            ),
            branch=branch,
        )  # fmt: skip
    _rows, counts = S.day_summary(branch, today)
    hr = HRSettings.for_branch(branch)
    return render(
        request,
        "hr/index.html",
        {
            "employees": qs.order_by("staff_type", "full_name")[:500],
            "counts": counts,
            "status": status,
            "staff_type": staff_type,
            "q": q,
            "statuses": Employee.Status.choices,
            "types": Employee.StaffType.choices,
            "settings_missing": hr.pk is None or hr.duty_starts is None,
            "can_edit": request.user.has_branch_perm(S.ADD_PERM, branch.id),
            "can_import": request.user.has_branch_perm(S.IMPORT_PERM, branch.id),
            "today": today,
        },
    )


@login_required
def employee_form(request, employee_id=None):
    branch = _branch(request, S.EDIT_PERM if employee_id else S.ADD_PERM)
    instance = _employee(branch, employee_id) if employee_id else None
    form = EmployeeForm(request.POST or None, request.FILES or None, instance=instance, branch=branch)
    if request.method == "POST" and form.is_valid():
        e = form.save()
        if instance and instance.status != Employee.Status.ACTIVE and e.left_on:
            S.rebuild_days(branch, e.left_on, timezone.localdate(), [e.id])
        messages.success(request, _("Saved."))
        return redirect("hr:employee", employee_id=e.id)
    return render(request, "hr/employee_form.html", {"form": form, "employee": instance})


@login_required
def employee_cv(request, employee_id):
    """The CV is private: only people who may see this branch's staff can open it."""
    branch = _branch(request)
    e = _employee(branch, employee_id)
    if not e.cv:
        raise Http404
    return FileResponse(e.cv.open("rb"), filename=f"CV - {e.full_name}{Path(e.cv.name).suffix}")


@login_required
def employee(request, employee_id):
    branch = _branch(request)
    e = _employee(branch, employee_id)
    month = _month(request)
    days = StaffDay.objects.filter(employee=e, date__gte=month, date__lte=_month_end(month))
    session = _session()
    return render(
        request,
        "hr/employee.html",
        {
            "employee": e,
            "month": month,
            "days": days,
            "balance": S.leave_balance(e, session),
            "leaves": e.leave_requests.select_related("leave_type").order_by("-starts_on")[:20],
            "can_edit": request.user.has_branch_perm(S.EDIT_PERM, branch.id),
            "can_leave": request.user.has_branch_perm(S.LEAVE_PERM, branch.id),
        },
    )


# --- Staff list from Excel -------------------------------------------------------------------


@login_required
def staff_import(request):
    from apps.people.views import _import_dir

    branch = _branch(request, S.ADD_PERM)
    if request.GET.get("template"):
        response = HttpResponse(
            importer.template_workbook(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="TASEMS-staff-template.xlsx"'
        return response
    context = {}
    if request.method == "POST" and request.POST.get("step") == "check":
        upload = request.FILES.get("file")
        if upload is None or not upload.name.lower().endswith(".xlsx") or upload.size > 5 * 1024 * 1024:
            messages.error(request, _("Choose an Excel file (.xlsx) smaller than 5 MB."))
        else:
            content = upload.read()
            result = importer.check(content, branch=branch)
            token = ""
            if result.ok or result.can_skip:
                token = secrets.token_hex(16)
                (_import_dir() / f"{token}.xlsx").write_bytes(content)
                request.session["staff_import"] = {"token": token, "branch": branch.id}
            context.update(result=result, token=token, shown_errors=result.errors[:300])
            if result.can_skip:
                context.update(skip_count=len(result.problem_rows), keep_count=len(result.without_problem_rows().rows))
    elif request.method == "POST" and request.POST.get("step") == "import":
        pending = request.session.pop("staff_import", None)
        token = request.POST.get("token", "")
        if not pending or pending["token"] != token or pending["branch"] != branch.id or not token.isalnum():
            messages.error(request, _("This import has expired. Please check the file again."))
            return redirect("hr:staff_import")
        path = _import_dir() / f"{token}.xlsx"
        try:
            content = path.read_bytes()
        except FileNotFoundError:
            messages.error(request, _("This import has expired. Please check the file again."))
            return redirect("hr:staff_import")
        finally:
            path.unlink(missing_ok=True)
        result = importer.check(content, branch=branch)
        skipped = []
        if request.POST.get("skip") == "1" and result.can_skip:
            skipped = result.problem_rows
            result = result.without_problem_rows()
        if not result.ok:
            messages.error(request, _("The data changed since the check. Please check the file again."))
            return redirect("hr:staff_import")
        n = importer.run_import(result, branch=branch)
        messages.success(request, _("%(n)s staff members added.") % {"n": n})
        if skipped:
            from apps.people.views import skipped_message

            messages.warning(request, skipped_message(skipped))
        return redirect("hr:index")
    return render(request, "hr/staff_import.html", context)


# --- Attendance machine file -----------------------------------------------------------------


@login_required
def punches_import(request):
    from apps.people.views import _import_dir

    branch = _branch(request, S.IMPORT_PERM)
    context = {}
    if request.method == "POST" and request.POST.get("step") == "check":
        upload = request.FILES.get("file")
        name = upload.name if upload else ""
        if upload is None or not name.lower().endswith((".xlsx", ".csv", ".dat", ".txt")):
            messages.error(request, _("Choose the machine's export file (.xlsx, .csv, .dat or .txt)."))
        elif upload.size > S.MAX_BYTES:
            messages.error(request, _("The file is larger than 10 MB. Export one month at a time."))
        else:
            content = upload.read()
            machine_file = S.read_machine_file(content, name)
            check = S.check_punches(branch, machine_file)
            token = ""
            if check.matched and not machine_file.errors:
                token = secrets.token_hex(16)
                ext = name.lower().rsplit(".", 1)[-1]
                (_import_dir() / f"{token}.xlsx").write_bytes(content)  # removed after a day like other uploads
                request.session["punch_import"] = {"token": token, "branch": branch.id, "name": name, "ext": ext}
            context.update(file=machine_file, check=check, token=token, filename=name)
    elif request.method == "POST" and request.POST.get("step") == "import":
        pending = request.session.pop("punch_import", None)
        token = request.POST.get("token", "")
        if not pending or pending["token"] != token or pending["branch"] != branch.id or not token.isalnum():
            messages.error(request, _("This import has expired. Please check the file again."))
            return redirect("hr:punches_import")
        path = _import_dir() / f"{token}.xlsx"
        try:
            content = path.read_bytes()
        except FileNotFoundError:
            messages.error(request, _("This import has expired. Please check the file again."))
            return redirect("hr:punches_import")
        finally:
            path.unlink(missing_ok=True)
        check = S.check_punches(branch, S.read_machine_file(content, pending["name"]))
        added = S.save_punches(branch, check, request.user, pending["name"])
        messages.success(
            request,
            _("%(n)s new punches saved for %(e)s staff; attendance worked out from %(a)s to %(b)s.")
            % {"n": added, "e": check.employees, "a": check.first, "b": check.last},
        )
        return redirect(f"{reverse('hr:day')}?date={check.last:%Y-%m-%d}")
    return render(request, "hr/punches_import.html", context)


# --- Daily and monthly attendance ------------------------------------------------------------


@login_required
def day(request):
    branch = _branch(request)
    today = timezone.localdate()
    date = min(_date(request.GET.get("date"), today), today)
    can_correct = request.user.has_branch_perm(S.CORRECT_PERM, branch.id)
    if request.method == "POST":
        if not can_correct:
            raise PermissionDenied
        e = _employee(branch, request.POST.get("employee"))
        if request.POST.get("remove"):
            S.remove_correction(e, date, request.user)
            messages.success(request, _("Correction removed for %(n)s.") % {"n": e.full_name})
        else:
            form = CorrectionForm(request.POST)
            if form.is_valid():
                d = form.cleaned_data
                S.correct_day(e, date, d["status"], d["note"], request.user, d["first_in"], d["last_out"])
                messages.success(request, _("Attendance corrected for %(n)s.") % {"n": e.full_name})
            else:
                messages.error(request, _("Choose a status and write the reason."))
        return redirect(f"{reverse('hr:day')}?date={date:%Y-%m-%d}")
    if request.GET.get("refresh") and can_correct:
        S.rebuild_days(branch, date, date)
    rows, counts = S.day_summary(branch, date)
    return render(
        request,
        "hr/day.html",
        {
            "date": date,
            "rows": rows,
            "counts": counts,
            "previous": date - dt.timedelta(days=1),
            "next": date + dt.timedelta(days=1) if date < today else None,
            "can_correct": can_correct,
            "statuses": [c for c in CorrectionForm.base_fields["status"].choices],
        },
    )


@login_required
def month(request):
    branch = _branch(request)
    first = _month(request)
    dates, rows = S.month_sheet(branch, first)
    if request.GET.get("format") == "xlsx":
        return _month_xlsx(branch, first, dates, rows)
    return render(request, "hr/month.html", {"month": first, "dates": dates, "rows": rows})


TOTAL_COLUMNS = [
    ("present", "P"),
    ("late", "L"),
    ("half_day", "H"),
    ("absent", "A"),
    ("leave", "LV"),
    ("unpaid", "Unpaid days"),
]


def _month_xlsx(branch, first, dates, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = f"{first:%b %Y}"
    ws.append([f"{branch.name} — staff attendance {first:%B %Y}"])
    ws.append(["Code", "Name", "Designation"] + [d.day for d in dates] + [c for _s, c in TOTAL_COLUMNS])
    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E3EFEB")
    for r in rows:
        e = r["employee"]
        ws.append(
            [e.code, e.full_name, e.designation]
            + [c.code if c else "" for c in r["cells"]]
            + [r["totals"].get(s, 0) for s, _c in TOTAL_COLUMNS]
        )
    ws.append([])
    ws.append(["P = present, L = late, H = half day, A = absent, LV = leave, HD = holiday, – = off day"])
    ws.append(["Unpaid days = absent days + days of unpaid leave (salary is deducted for these)."])
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 18
    ws.freeze_panes = "D3"
    buffer = io.BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="Staff-attendance-{first:%Y-%m}.xlsx"'
    return response


# --- Leave -----------------------------------------------------------------------------------


@login_required
def leaves(request):
    branch = _branch(request)
    can_decide = request.user.has_branch_perm(S.APPROVE_PERM, branch.id)
    if request.method == "POST":
        if not can_decide:
            raise PermissionDenied
        req = LeaveRequest.objects.filter(pk=request.POST.get("id"), branch=branch).select_related("employee").first()
        if req is None:
            raise Http404
        try:
            if request.POST.get("decision") == "cancel":
                S.cancel_leave(req, request.user)
            else:
                S.decide_leave(
                    req, request.user, request.POST.get("decision") == "approve", request.POST.get("note", "")
                )
            messages.success(request, _("Decision saved."))
        except ValueError as exc:
            messages.error(request, str(exc))
        return redirect("hr:leaves")
    qs = LeaveRequest.objects.filter(branch=branch).select_related(
        "employee", "leave_type", "requested_by", "decided_by"
    )
    return render(
        request,
        "hr/leaves.html",
        {
            "pending": qs.filter(status=LeaveRequest.Status.PENDING).order_by("starts_on"),
            "recent": qs.exclude(status=LeaveRequest.Status.PENDING).order_by("-starts_on")[:100],
            "can_decide": can_decide,
        },
    )


def _leave_form(request, employee, back):
    form = LeaveForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        d = form.cleaned_data
        try:
            S.request_leave(employee, d["leave_type"], d["starts_on"], d["ends_on"], d["reason"], request.user)
            messages.success(request, _("Leave request saved. It counts once the principal approves it."))
            return redirect(back)
        except ValueError as exc:
            form.add_error(None, str(exc))
    return render(request, "hr/leave_form.html", {"form": form, "employee": employee, "back": back})


@login_required
def leave_new(request, employee_id):
    branch = _branch(request, S.LEAVE_PERM)
    e = _employee(branch, employee_id)
    return _leave_form(request, e, reverse("hr:employee", args=[e.id]))


# --- Self-service ----------------------------------------------------------------------------


def _me(request):
    e = Employee.objects.filter(user=request.user).select_related("branch").first()
    if e is None:
        raise Http404
    return e


@login_required
def mine(request):
    e = _me(request)
    first = _month(request)
    return render(
        request,
        "hr/mine.html",
        {
            "employee": e,
            "month": first,
            "days": StaffDay.objects.filter(employee=e, date__gte=first, date__lte=_month_end(first)),
            "balance": S.leave_balance(e, _session()),
            "leaves": e.leave_requests.select_related("leave_type").order_by("-starts_on")[:20],
        },
    )


@login_required
def my_leave(request):
    e = _me(request)
    return _leave_form(request, e, reverse("hr:mine"))


@login_required
@require_POST
def my_leave_cancel(request, leave_id):
    e = _me(request)
    req = LeaveRequest.objects.filter(pk=leave_id, employee=e, status=LeaveRequest.Status.PENDING).first()
    if req is None:
        raise Http404
    S.cancel_leave(req, request.user)
    messages.success(request, _("Leave request withdrawn."))
    return redirect("hr:mine")
