"""Staff list from Excel: checked row by row, saved all together or not at all."""

import datetime as dt
import io
import re
from dataclasses import dataclass, field

from django.db import transaction
from django.utils.translation import gettext as _
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from apps.core.phone import first_mobile, normalize_cnic

from .models import Employee

MAX_ROWS = 2000

COLUMNS = [
    ("code", "Employee Code", True),
    ("full_name", "Full Name", True),
    ("father_name", "Father Name", False),
    ("gender", "Gender", True),
    ("cnic", "CNIC", False),
    ("phone", "Mobile", False),
    ("designation", "Designation", True),
    ("staff_type", "Staff Type", True),
    ("joined_on", "Joining Date", True),
    ("machine_id", "Machine ID", False),
    ("qualification", "Qualification", False),
]
HEADER_TO_KEY = {re.sub(r"[^a-z]", "", h.lower()): k for k, h, _r in COLUMNS}
GENDERS = {"male": "male", "m": "male", "man": "male", "female": "female", "f": "female", "woman": "female"}
STAFF_TYPES = {
    "teacher": "teaching",
    "teaching": "teaching",
    "hifz": "hifz",
    "hifzteacher": "hifz",
    "qari": "hifz",
    "admin": "admin",
    "office": "admin",
    "administration": "admin",
    "support": "support",
    "supportstaff": "support",
}


def _key(v):
    return re.sub(r"[^a-z]", "", str(v or "").lower())


def _text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _date(v):
    if v in (None, ""):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%b-%Y"):
        try:
            return dt.datetime.strptime(str(v).strip(), fmt).date()
        except ValueError:
            continue
    raise ValueError


def template_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Staff"
    ws.append([h + (" *" if r else "") for _k, h, r in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E3EFEB")
    for i in range(len(COLUMNS)):
        ws.column_dimensions[chr(65 + i)].width = 18
    ws.freeze_panes = "A2"
    guide = wb.create_sheet("How to fill")
    for line in [
        ("Columns marked * are required. Keep the header row as it is.",),
        ("Gender", "Male or Female"),
        ("Staff Type", "Teacher, Hifz (or Qari), Admin (or Office), Support"),
        ("Joining Date", "Day first: 01/04/2020, or an Excel date"),
        ("Machine ID", "The number the person is enrolled under on the face-recognition machine"),
        ("CNIC / Mobile", "33100-1234567-1 / 0300-1234567"),
    ]:
        guide.append(list(line))
    guide.column_dimensions["A"].width = 16
    guide.column_dimensions["B"].width = 80
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@dataclass
class Result:
    rows: list = field(default_factory=list)
    errors: list = field(default_factory=list)  # (row, column, message)
    numbers: list = field(default_factory=list)  # the Excel row number of each entry in rows
    warnings: list = field(default_factory=list)  # imported anyway, as written or left empty

    @property
    def ok(self):
        return not self.errors and bool(self.rows)

    @property
    def problem_rows(self):
        return sorted({n for n, _c, _m in self.errors if n is not None})

    @property
    def can_skip(self):
        bad = set(self.problem_rows)
        return (
            bool(self.errors)
            and all(n is not None for n, _c, _m in self.errors)
            and any(n not in bad for n in self.numbers)
        )

    def without_problem_rows(self):
        bad = set(self.problem_rows)
        keep = [(d, n) for d, n in zip(self.rows, self.numbers, strict=True) if n not in bad]
        return Result(
            rows=[d for d, _n in keep],
            numbers=[n for _d, n in keep],
            warnings=[w for w in self.warnings if w[0] not in bad],
        )


def check(content, *, branch):
    result = Result()
    try:
        ws = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
    except Exception:  # noqa: BLE001
        result.errors.append(
            (None, "", _("This file could not be opened. Save it as an Excel .xlsx file and try again."))
        )
        return result
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None) or ()
    positions = {}
    for i, title in enumerate(header):
        key = HEADER_TO_KEY.get(_key(title))
        if key:
            positions[key] = i
    missing = [h for k, h, r in COLUMNS if r and k not in positions]
    if missing:
        result.errors.append((None, "", _("Missing columns: %(c)s") % {"c": ", ".join(missing)}))
        return result
    headers = {k: h for k, h, _r in COLUMNS}
    codes = {c.lower() for c in Employee.objects.filter(branch=branch).values_list("code", flat=True)}
    machines = set(Employee.objects.filter(branch=branch).exclude(machine_id="").values_list("machine_id", flat=True))
    seen_codes, seen_machines = {}, {}
    for number, values in enumerate(rows, start=2):
        if values is None or all(v in (None, "") for v in values):
            continue
        if len(result.rows) >= MAX_ROWS:
            result.errors.append((None, "", _("Too many rows.")))
            break
        raw = {k: (values[i] if i < len(values) else None) for k, i in positions.items()}
        d = {k: _text(raw.get(k)) for k, _h, _r in COLUMNS if k != "joined_on"}

        def err(col, msg, _n=number):
            result.errors.append((_n, headers[col], msg))

        def warn(col, msg, _n=number):
            result.warnings.append((_n, headers[col], msg))

        for k, _h, required in COLUMNS:
            if required and k != "joined_on" and not d[k]:
                err(k, _("Required"))
        if d["code"]:
            if d["code"].lower() in codes:
                err("code", _("Already exists in this branch"))
            elif d["code"].lower() in seen_codes:
                err("code", _("Repeated in this file (row %(r)s)") % {"r": seen_codes[d["code"].lower()]})
            seen_codes.setdefault(d["code"].lower(), number)
        if d["machine_id"]:
            if d["machine_id"] in machines:
                err("machine_id", _("Another person already has this machine ID"))
            elif d["machine_id"] in seen_machines:
                err("machine_id", _("Repeated in this file (row %(r)s)") % {"r": seen_machines[d["machine_id"]]})
            seen_machines.setdefault(d["machine_id"], number)
        if d["gender"]:
            d["gender"] = GENDERS.get(d["gender"].lower())
            if d["gender"] is None:
                err("gender", _("Use Male or Female"))
        if d["staff_type"]:
            d["staff_type"] = STAFF_TYPES.get(_key(d["staff_type"]))
            if d["staff_type"] is None:
                err("staff_type", _("Use Teacher, Hifz, Admin or Support"))
        if d["cnic"]:
            given = str(d["cnic"])[:30]
            d["cnic"] = normalize_cnic(d["cnic"]) or ""
            if not d["cnic"]:
                warn("cnic", _("A CNIC has 13 digits: «%(value)s». Left empty") % {"value": given})
        if d["phone"]:
            given = re.sub(r"\s+", " ", str(d["phone"])).strip()[:20]
            d["phone"] = first_mobile(d["phone"]) or ""
            if not d["phone"]:
                # Imported as written, in "other phone" (the mobile field holds 11 digits only).
                d["alt_phone"] = given
                warn("phone", _("Not a Pakistani mobile number: «%(value)s». Imported as written") % {"value": given})
        try:
            d["joined_on"] = _date(raw.get("joined_on"))
            if d["joined_on"] is None:
                err("joined_on", _("Required"))
        except ValueError:
            d["joined_on"] = None
            err("joined_on", _("Not a date. Use day/month/year, e.g. 01/04/2020"))
        result.rows.append(d)
        result.numbers.append(number)
    if not result.rows and not result.errors:
        result.errors.append((None, "", _("The file has no staff rows.")))
    return result


@transaction.atomic
def run_import(result, *, branch):
    if not result.ok:
        raise ValueError("Import called on a file with problems")
    for d in result.rows:
        e = Employee(branch=branch, **d)
        e.full_clean()
        e.save()
    return len(result.rows)
