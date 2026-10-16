"""Files exchanged with payment partners.

JazzCash is not connected by API (confirmed 16 Oct 2026). Each month the school gives JazzCash a
file of Voucher IDs and amounts; JazzCash sends back a file of what was paid against which
Voucher ID, and the school records those payments. The same import works for any list of
payments by Voucher ID (e.g. a bank statement). Bank Alfalah is connected by API to the old
software; moving that connection to TASEMS needs Bank Alfalah's interface document.
"""

import datetime as dt
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils.translation import gettext as _
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .models import Payment, Voucher
from .services import ZERO, record_payment

MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024

EXPORT_COLUMNS = [
    ("Voucher ID", 14),
    ("Student Name", 24),
    ("Father Name", 24),
    ("Class", 14),
    ("Section", 10),
    ("Admission No", 13),
    ("Fee Month", 16),
    ("Due Date", 12),
    ("Amount Payable by Due Date", 16),
    ("Amount Payable after Due Date", 16),
]


def open_vouchers(branch, month):
    return (
        Voucher.objects.filter(
            branch=branch, first_month__lte=month, last_month__gte=month, status=Voucher.Status.OPEN, payable__gt=0
        )
        .select_related("student", "section", "section__class_level")
        .order_by("number")
    )


def export_open_vouchers(branch, month):
    """Excel list of unpaid vouchers for a month, to give to JazzCash."""
    wb = Workbook()
    ws = wb.active
    ws.title = f"Vouchers {month:%b %Y}"
    ws.append([h for h, _w in EXPORT_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E3EFEB")
    for i, (_h, width) in enumerate(EXPORT_COLUMNS):
        ws.column_dimensions[chr(65 + i)].width = width
    ws.freeze_panes = "A2"
    count = 0
    for v in open_vouchers(branch, month):
        months = f"{v.first_month:%b %Y}" + (f" – {v.last_month:%b %Y}" if v.last_month != v.first_month else "")
        ws.append(
            [
                v.number,
                v.student.full_name,
                v.student.father_name,
                v.section.class_level.name if v.section else "",
                v.section.name if v.section else "",
                v.student.admission_no,
                months,
                v.due_on,
                int(v.payable) if v.payable == v.payable.to_integral_value() else float(v.payable),
                float(v.payable_after_due),
            ]
        )
        ws.cell(row=ws.max_row, column=8).number_format = "DD/MM/YYYY"
        count += 1
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), count


# --- Paid file -------------------------------------------------------------------------------

HEADERS = {
    "voucher": {
        "voucherid",
        "voucherno",
        "vouchernumber",
        "voucher",
        "billno",
        "billnumber",
        "consumerno",
        "consumernumber",
        "challanno",
        "referenceno",
        "invoiceno",
    },
    "amount": {"amount", "paidamount", "amountpaid", "paid", "netamount", "transactionamount", "totalamount"},
    "date": {"date", "paiddate", "paymentdate", "transactiondate", "datepaid", "txndate", "valuedate"},
    "reference": {"transactionid", "tid", "txnid", "transactionno", "transactionref", "reference", "stan", "rrn"},
}


def _key(text):
    return re.sub(r"[^a-z]", "", str(text or "").lower())


def _date(value):
    if value in (None, ""):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value).strip().split(" ")[0]
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%b-%Y", "%d-%b-%y"):
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError


@dataclass
class PaidRow:
    number: int  # row in the file
    voucher_no: str
    amount: Decimal | None
    paid_on: dt.date | None
    reference: str
    voucher: Voucher | None = None
    problem: str = ""

    @property
    def ok(self):
        return not self.problem


@dataclass
class PaidFile:
    rows: list = field(default_factory=list)
    errors: list = field(default_factory=list)

    @property
    def good(self):
        return [r for r in self.rows if r.ok]

    @property
    def total(self):
        return sum((r.amount for r in self.good), ZERO)


def read_paid_file(content, *, branch, method, default_date, today):
    """Reads an Excel file of payments by Voucher ID and checks every row. Nothing is saved."""
    result = PaidFile()
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - any unreadable file gets the same message
        result.errors.append(_("This file could not be opened. Save it as an Excel .xlsx file and try again."))
        return result
    rows = wb.worksheets[0].iter_rows(values_only=True)
    positions, header_row = {}, 0
    # The header may sit below a title line or two.
    for header_row, values in enumerate(rows, start=1):
        for i, title in enumerate(values or ()):
            for key, names in HEADERS.items():
                if _key(title) in names and key not in positions:
                    positions[key] = i
        if "voucher" in positions and "amount" in positions:
            break
        positions = {}
        if header_row >= 10:
            break
    if "voucher" not in positions or "amount" not in positions:
        result.errors.append(_("Could not find the 'Voucher ID' and 'Amount' columns in the first rows of the file."))
        return result

    seen = {}
    for number, values in enumerate(rows, start=header_row + 1):
        if values is None or all(v in (None, "") for v in values):
            continue
        if len(result.rows) >= MAX_ROWS:
            result.errors.append(_("Too many rows. Import at most %(n)s payments at a time.") % {"n": MAX_ROWS})
            break

        def cell(key, _values=values):
            i = positions.get(key)
            return _values[i] if i is not None and i < len(_values) else None

        raw_no = cell("voucher")
        voucher_no = (
            str(int(raw_no)) if isinstance(raw_no, float) and raw_no.is_integer() else str(raw_no or "").strip()
        )
        row = PaidRow(
            number=number,
            voucher_no=voucher_no,
            amount=None,
            paid_on=None,
            reference=str(cell("reference") or "").strip(),
        )
        try:
            amount = Decimal(str(cell("amount")).replace(",", "").strip())
            row.amount = amount if amount > 0 else None
        except (InvalidOperation, ValueError):
            row.amount = None
        try:
            row.paid_on = _date(cell("date")) or default_date
        except ValueError:
            row.problem = _("Date not understood")
        if not row.problem and row.amount is None:
            row.problem = _("Amount missing or not a number")
        if not row.problem and row.paid_on and row.paid_on > today:
            row.problem = _("Date is in the future")
        if not row.problem:
            digits = re.sub(r"\D", "", voucher_no)
            row.voucher = (
                Voucher.objects.filter(branch=branch, number=int(digits)).select_related("student").first()
                if digits
                else None
            )
            if row.voucher is None:
                row.problem = _("No voucher with this ID in this branch")
            elif row.voucher.status == Voucher.Status.CANCELLED:
                row.problem = _("Voucher was cancelled")
        if not row.problem:
            key = (row.voucher.pk, row.amount, row.paid_on, row.reference)
            if key in seen:
                row.problem = _("Same payment repeated in this file (row %(r)s)") % {"r": seen[key]}
            seen.setdefault(key, number)
        if not row.problem and already_recorded(row, method):
            row.problem = _("Already recorded")
        result.rows.append(row)
    if not result.rows and not result.errors:
        result.errors.append(_("The file has no payment rows."))
    return result


def already_recorded(row, method):
    qs = Payment.objects.filter(voucher=row.voucher, method=method, is_reversed=False)
    if row.reference:
        return qs.filter(reference=row.reference).exists()
    return qs.filter(amount=row.amount, paid_on=row.paid_on).exists()


@transaction.atomic
def record_paid_rows(rows, *, method, user, note):
    """Records the good rows. Each payment goes through the normal rules (late fine, allocation)."""
    payments = []
    for r in rows:
        if not r.ok:
            continue
        payments.append(
            record_payment(
                r.voucher.student,
                r.amount,
                r.paid_on,
                method,
                user,
                voucher=r.voucher,
                reference=r.reference,
                note=note[:200],
            )
        )
    return payments
