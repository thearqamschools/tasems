"""Fees: the rules confirmed for the Samanabad campus on 16 Oct 2026."""

import datetime as dt
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone

from apps.academics.models import ClassLevel, Section
from apps.accounts.models import Scope
from apps.fees import reports
from apps.fees import services as S
from apps.fees.models import BankAccount, BranchFeeSettings, Charge, Concession, FeeHead, FeeRate, Payment, Voucher
from apps.people.models import Enrolment, Student

pytestmark = pytest.mark.django_db
TODAY = {"date": dt.date(2026, 10, 14)}  # a Wednesday, after the 10th
OCT, SEP, NOV, DEC = (dt.date(2026, m, 1) for m in (10, 9, 11, 12))
D = Decimal


@pytest.fixture(autouse=True)
def fixed_today(monkeypatch):
    monkeypatch.setattr(timezone, "localdate", lambda *args, **kwargs: TODAY["date"])
    yield
    TODAY["date"] = dt.date(2026, 10, 14)


@pytest.fixture
def fees(org, make_user):
    branch = org["samanabad"]
    accountant = make_user("accountant", Scope.BRANCH, branch)
    sections = {}

    def section(code, name="A"):
        key = (code, name)
        if key not in sections:
            level = ClassLevel.objects.get(code=code)
            sections[key] = Section.objects.create(
                branch=branch,
                session=org["session"],
                class_level=level,
                name=name,
                gender_group="boys" if level.separate_by_gender else "mixed",
            )
        return sections[key]

    counter = {"n": 0}

    def student(code="class-4", cnic="", admitted=dt.date(2026, 4, 1), branch_=branch):
        counter["n"] += 1
        s = Student.objects.create(
            branch=branch_,
            admission_no=f"FEE-{counter['n']}",
            first_name=f"Child{counter['n']}",
            gender="male",
            admission_date=admitted,
            father_cnic=cnic,
            father_name="Test Father",
        )
        if branch_ == branch:
            Enrolment.objects.create(branch=branch, student=s, session=org["session"], section=section(code))
        return s

    return {"branch": branch, "user": accountant, "student": student, "session": org["session"], "section": section}


def generate(fees, first=OCT, last=None, due=None, sections=None):
    return S.generate_vouchers(
        fees["branch"],
        fees["session"],
        first,
        last or first,
        due or S.default_due_date(fees["branch"], first),
        fees["user"],
        sections,
    )


def voucher_of(student):
    return Voucher.objects.filter(student=student).exclude(status="cancelled").order_by("-pk").first()


def amounts(voucher):
    return {c.head.code: c.amount for c in voucher.charges.filter(is_cancelled=False).select_related("head")}


def test_heads_and_samanabad_rates_are_seeded(seeded):
    assert list(FeeHead.objects.values_list("short_name", flat=True)) == [
        "Admission",
        "Reg.",
        "Tuition",
        "Annual",
        "S.V.W.",
        "Gen. Exp.",
        "Others",
        "Fine",
    ]
    rate = FeeRate.objects.get(branch__code="fsd-samanabad", head__code="tuition")
    assert (rate.amount, rate.sibling_amount) == (D(4300), D(4000))
    settings = BranchFeeSettings.objects.get(branch__code="fsd-samanabad")
    assert (settings.due_day, settings.late_fine, settings.account_number) == (10, D(200), "")


def test_due_date_is_the_10th_or_the_day_made(fees):
    TODAY["date"] = dt.date(2026, 10, 3)
    assert S.default_due_date(fees["branch"], OCT) == dt.date(2026, 10, 10)
    TODAY["date"] = dt.date(2026, 10, 14)
    assert S.default_due_date(fees["branch"], OCT) == dt.date(2026, 10, 14)
    assert S.default_due_date(fees["branch"], NOV) == dt.date(2026, 11, 10)  # made in advance


def test_monthly_voucher_and_sibling_fee(fees):
    elder = fees["student"](cnic="3310012345671", admitted=dt.date(2024, 4, 1))
    younger = fees["student"](cnic="3310012345671", admitted=dt.date(2025, 4, 1))
    other = fees["student"](cnic="3310099999991")
    result = generate(fees)
    assert len(result["created"]) == 3
    assert amounts(voucher_of(elder)) == {"tuition": D(4300)}
    assert amounts(voucher_of(younger)) == {"tuition": D(4000)}
    assert amounts(voucher_of(other)) == {"tuition": D(4300)}
    assert voucher_of(younger).charges.get().discount_reason == "Sibling"


def test_sibling_in_another_branch_does_not_count(fees, org):
    fees["student"](cnic="3310012345671", branch_=org["fsd_franchise"], admitted=dt.date(2020, 1, 1))
    child = fees["student"](cnic="3310012345671")
    generate(fees)
    assert amounts(voucher_of(child)) == {"tuition": D(4300)}


def test_annual_charges_in_april_by_class(fees):
    primary, pre9, hifz = fees["student"]("class-3"), fees["student"]("pre-9"), fees["student"]("hifz")
    generate(fees, first=dt.date(2026, 4, 1))
    assert amounts(voucher_of(primary)) == {"tuition": D(4300), "annual": D(4500)}
    assert amounts(voucher_of(pre9))["annual"] == D(5500)
    assert amounts(voucher_of(hifz))["annual"] == D(2000)


def test_generating_twice_skips_students_already_billed(fees):
    fees["student"]()
    generate(fees)
    again = generate(fees)
    assert again["created"] == [] and again["skipped"] == 1
    assert Charge.objects.count() == 1


def test_voucher_numbers_are_unique_and_sequential(fees):
    a, b = fees["student"](), fees["student"]()
    generate(fees)
    assert voucher_of(b).number == voucher_of(a).number + 1


def test_unpaid_month_moves_forward_with_a_fine(fees):
    child = fees["student"]()
    TODAY["date"] = dt.date(2026, 9, 1)
    generate(fees, first=SEP)
    sep = voucher_of(child)
    assert sep.due_on == dt.date(2026, 9, 10) and sep.payable == D(4300) and sep.late_fine == D(200)

    TODAY["date"] = dt.date(2026, 10, 1)
    generate(fees, first=OCT)
    oct_ = voucher_of(child)
    sep.refresh_from_db()
    assert sep.status == Voucher.Status.SUPERSEDED
    assert amounts(oct_) == {"tuition": D(4300), "fine": D(200)}
    assert oct_.arrears == D(4300)
    assert oct_.payable == D(8800)  # September 4300 + fine 200 + October 4300
    table = reports.voucher_table(oct_)
    assert table.rows[0][0] == "Previous dues" and table.total == D(8800)


def test_on_time_payment_has_no_fine(fees):
    child = fees["student"]()
    TODAY["date"] = dt.date(2026, 9, 1)
    generate(fees, first=SEP)
    sep = voucher_of(child)
    S.record_payment(child, D(4300), dt.date(2026, 9, 10), "bank_branch", fees["user"], voucher=sep)
    sep.refresh_from_db()
    assert sep.status == Voucher.Status.PAID
    TODAY["date"] = dt.date(2026, 10, 1)
    generate(fees, first=OCT)
    assert amounts(voucher_of(child)) == {"tuition": D(4300)} and voucher_of(child).arrears == 0


def test_late_payment_brings_the_fine(fees):
    child = fees["student"]()
    TODAY["date"] = dt.date(2026, 9, 1)
    generate(fees, first=SEP)
    sep = voucher_of(child)
    TODAY["date"] = dt.date(2026, 9, 15)
    due, fine = S.amount_due_now(sep)
    assert (due, fine) == (D(4500), D(200))
    S.record_payment(child, D(4500), TODAY["date"], "counter", fees["user"], voucher=sep)
    sep.refresh_from_db()
    assert sep.status == Voucher.Status.PAID and S.balance(child) == 0
    assert Charge.objects.filter(student=child, head__code="fine").count() == 1


def test_partial_payment_keeps_voucher_open(fees):
    child = fees["student"]()
    generate(fees)
    v = voucher_of(child)
    S.record_payment(child, D(2000), TODAY["date"], "counter", fees["user"], voucher=v)
    v.refresh_from_db()
    assert v.status == Voucher.Status.OPEN and S.balance(child) == D(2300)
    S.record_payment(child, D(2300), TODAY["date"], "jazzcash", fees["user"], voucher=v, reference="JC123")
    v.refresh_from_db()
    assert v.status == Voucher.Status.PAID


def test_several_months_together(fees):
    child = fees["student"]()
    v = S.issue_voucher(child, fees["session"], OCT, DEC, dt.date(2026, 10, 20), fees["user"])
    assert v.payable == D(12900) and len(v.months) == 3
    assert generate(fees, first=NOV)["skipped"] == 1  # November already billed


def test_admission_voucher_with_one_time_heads(fees):
    child = fees["student"]("play-group")
    heads = {h.code: h for h in FeeHead.objects.all()}
    level = child.enrolments.get().section.class_level_id
    extras = [
        (heads[c], S.one_time_price(child, heads[c], level, OCT), D(0), "", OCT, "At admission")
        for c in ("admission", "registration", "annual")
    ]
    v = S.issue_voucher(child, fees["session"], OCT, OCT, dt.date(2026, 10, 20), fees["user"], extras)
    assert amounts(v) == {"admission": D(6500), "registration": D(1500), "annual": D(4500), "tuition": D(4300)}


def test_special_discount_needs_approval(fees, make_user):
    child = fees["student"](cnic="3310012345671")
    fees["student"](cnic="3310012345671", admitted=dt.date(2020, 1, 1))  # makes `child` a sibling (4000)
    tuition = FeeHead.objects.get(code="tuition")
    c = Concession.objects.create(
        branch=fees["branch"],
        student=child,
        head=tuition,
        kind="percent",
        value=50,
        starts_on=OCT,
        reason="Staff child",
        requested_by=fees["user"],
    )
    generate(fees)
    assert amounts(voucher_of(child)) == {"tuition": D(4000)}  # pending: sibling fee only
    S.decide_concession(c, make_user("principal", Scope.BRANCH, fees["branch"]), approve=True)
    generate(fees, first=NOV)
    charge = voucher_of(child).charges.get(month=NOV)
    assert charge.amount == D(2150) and charge.discount_reason.startswith("Special discount")


def test_payments_pay_the_oldest_month_first(fees):
    child = fees["student"]()
    TODAY["date"] = dt.date(2026, 9, 1)
    generate(fees, first=SEP)
    TODAY["date"] = dt.date(2026, 10, 1)
    generate(fees, first=OCT)
    S.record_payment(child, D(4300), dt.date(2026, 10, 5), "counter", fees["user"])
    left = {(c.head.code, c.month): amt for c, amt in S.unpaid_charges(child)}
    assert left == {("fine", SEP): D(200), ("tuition", OCT): D(4300)}


def test_reversing_a_payment_reopens_the_voucher(fees):
    child = fees["student"]()
    generate(fees)
    v = voucher_of(child)
    p = S.record_payment(child, D(4300), TODAY["date"], "counter", fees["user"], voucher=v)
    S.reverse_payment(p, fees["user"], "Entered twice")
    v.refresh_from_db()
    assert v.status == Voucher.Status.OPEN and S.balance(child) == D(4300)


def test_cancelling_a_voucher_lets_it_be_made_again(fees):
    child = fees["student"]()
    generate(fees)
    S.cancel_voucher(voucher_of(child), fees["user"], "Wrong class")
    assert S.balance(child) == 0
    assert len(generate(fees)["created"]) == 1
    assert S.balance(child) == D(4300)


def test_voucher_with_payment_cannot_be_cancelled(fees):
    child = fees["student"]()
    generate(fees)
    v = voucher_of(child)
    S.record_payment(child, D(100), TODAY["date"], "counter", fees["user"], voucher=v)
    with pytest.raises(ValueError):
        S.cancel_voucher(v, fees["user"], "x")


def test_defaulters_and_month_summary(fees):
    payer, defaulter = fees["student"](), fees["student"]()
    TODAY["date"] = dt.date(2026, 9, 1)
    generate(fees, first=SEP)
    S.record_payment(payer, D(4300), dt.date(2026, 9, 5), "counter", fees["user"])
    TODAY["date"] = dt.date(2026, 10, 1)
    generate(fees, first=OCT)
    rows = {r["student_id"]: r for r in reports.outstanding(fees["branch"])}
    assert rows[defaulter.id]["due"] == D(8800) and rows[defaulter.id]["months"] == 2
    assert rows[defaulter.id]["oldest"] == SEP
    summary = reports.month_summary(fees["branch"], SEP)
    assert summary["collected"] == D(4300) and summary["billed"] == D(8600 + 200)


def test_royalty_on_all_fees_received(fees, org, make_user):
    franchise = org["fsd_franchise"]
    level = ClassLevel.objects.get(code="class-1")
    sec = Section.objects.create(branch=franchise, session=org["session"], class_level=level, name="A")
    FeeRate.objects.create(
        branch=franchise, head=FeeHead.objects.get(code="tuition"), amount=3000, effective_from=dt.date(2026, 4, 1)
    )
    kid = Student.objects.create(
        branch=franchise, admission_no="R-1", first_name="R", gender="male", admission_date=dt.date(2026, 4, 1)
    )
    Enrolment.objects.create(branch=franchise, student=kid, session=org["session"], section=sec)
    owner = make_user("principal", Scope.BRANCH, franchise)
    S.generate_vouchers(franchise, org["session"], OCT, OCT, dt.date(2026, 10, 20), owner)
    S.add_charge(kid, FeeHead.objects.get(code="others"), OCT, D(500), owner, note="Books")
    S.record_payment(kid, D(3500), dt.date(2026, 10, 15), "counter", owner)
    row = reports.royalty([franchise], OCT)[0]
    # 8% of everything received in the month (confirmed 16 Oct 2026)
    assert (row.collected, row.billed, row.percent, row.royalty) == (D(3500), D(3500), D(8), D(280))


def test_royalty_paid_pending_and_total_due_carry_forward(fees, org, make_user, login_as, tmp_path, settings):
    """The project lead's example: month 1 royalty 100,000, paid 50,000; month 2 royalty 20,000,
    paid 10,000 → total due 60,000."""
    settings.MEDIA_ROOT = tmp_path
    franchise = org["fsd_franchise"]
    franchise.royalty_percent = D(10)
    franchise.save()
    kid = Student.objects.create(
        branch=franchise, admission_no="R-1", first_name="R", gender="male", admission_date=dt.date(2026, 4, 1)
    )
    owner = make_user("principal", Scope.BRANCH, franchise)
    S.record_payment(kid, D(1_000_000), dt.date(2026, 9, 10), "counter", owner)  # Sep: royalty 100,000
    S.record_payment(kid, D(200_000), dt.date(2026, 10, 10), "counter", owner)  # Oct: royalty 20,000
    admin = make_user("head_office_admin", Scope.HEAD_OFFICE, org["fsd"])
    office = login_as(admin)
    url = f"/fees/royalty/branch/{franchise.id}/?month=2026-10"
    slip = SimpleUploadedFile("slip.jpg", b"\xff\xd8 slip", content_type="image/jpeg")
    office.post(url, {"month": "2026-09", "amount": "50000", "paid_on": "2026-09-30", "attachment": slip})
    office.post(url, {"month": "2026-10", "amount": "10000", "paid_on": "2026-10-14"})
    ledger = reports.royalty_ledger(franchise, OCT)
    sep, octo = ledger["rows"][-2], ledger["rows"][-1]
    assert (sep["due"], sep["paid"], sep["left"], sep["total_due"]) == (D(100_000), D(50_000), D(50_000), D(50_000))
    assert (octo["due"], octo["paid"], octo["left"], octo["total_due"]) == (D(20_000), D(10_000), D(10_000), D(60_000))
    assert reports.royalty_balances([franchise], OCT)[franchise.id] == {"paid": D(10_000), "total_due": D(60_000)}
    statement = office.get("/fees/royalty/?month=2026-10").content.decode()
    assert "60,000" in statement
    from apps.fees.models import RoyaltyPayment

    proof = RoyaltyPayment.objects.get(month=dt.date(2026, 9, 1))
    assert office.get(f"/fees/royalty/payment/{proof.id}/attachment/").status_code == 200
    # The branch sees its own royalty but cannot record payments; another branch cannot see it.
    branch_view = login_as(owner)
    assert "60,000" in branch_view.get(url).content.decode()
    assert branch_view.post(url, {"month": "2026-10", "amount": "1", "paid_on": "2026-10-14"}).status_code == 403
    outsider = login_as(make_user("principal", Scope.BRANCH, org["lhr_franchise"]))
    assert outsider.get(url).status_code == 404
    assert outsider.get(f"/fees/royalty/payment/{proof.id}/attachment/").status_code == 404


# --- Screens ---------------------------------------------------------------------------------


def test_accountant_makes_prints_and_receives(fees, login_as):
    child = fees["student"]()
    client = login_as(fees["user"])
    assert client.get("/fees/").status_code == 200
    form = {"first_month": "2026-10", "last_month": "2026-10", "due_on": "2026-10-14"}
    page = client.post("/fees/generate/", form).content.decode()
    assert "Make 1 voucher" in page and not Voucher.objects.exists()  # a check first, nothing saved
    client.post("/fees/generate/", {**form, "confirm": "1"})
    v = voucher_of(child)
    assert v.payable == D(4300)
    assert "Voucher" in client.get("/fees/vouchers/?month=2026-10").content.decode()
    pdf = client.get("/fees/vouchers/print/?month=2026-10")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")

    found = client.get(f"/fees/?q={v.number}")
    assert found.status_code == 302 and found.url == f"/fees/voucher/{v.id}/receive/"
    jazz = BankAccount.objects.create(branch=child.branch, name="JazzCash")
    bad = client.post(found.url, {"amount": "4300", "paid_on": "2026-10-14", "how": "bank"})
    assert "Choose the bank" in bad.content.decode()  # bank payments say which bank
    done = client.post(
        found.url, {"amount": "4300", "paid_on": "2026-10-14", "how": "bank", "bank": jazz.id, "reference": "J1"}
    )
    receipt = client.get(done.url).content.decode()
    assert "Fee receipt" in receipt and "4,300" in receipt and "JazzCash" in receipt
    v.refresh_from_db()
    assert v.status == Voucher.Status.PAID
    payment = Payment.objects.get(voucher=v)
    assert payment.method == "bank" and payment.bank == jazz and payment.received_by == fees["user"]


GENERATE = {"first_month": "2026-10", "last_month": "2026-10", "due_on": "2026-10-14", "confirm": "1"}


def test_barcode_page_receives_the_full_amount_and_shows_who(fees, login_as, tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path
    child = fees["student"]("class-4")
    client = login_as(fees["user"])
    client.post("/fees/generate/", GENERATE)
    v = voucher_of(child)
    assert (
        client.post("/fees/barcode/", {"voucher_number": "999999", "paid_on": "2026-10-14", "how": "cash"}).status_code
        == 200
    )
    proof = SimpleUploadedFile("slip.jpg", b"\xff\xd8 fake jpeg", content_type="image/jpeg")
    done = client.post(
        "/fees/barcode/",
        {"voucher_number": str(v.number), "paid_on": "2026-10-14", "how": "cash", "attachment": proof},
    )
    assert done.status_code == 302
    payment = Payment.objects.get(voucher=v)
    assert payment.amount == v.payable and payment.method == "counter" and payment.attachment
    again = client.post("/fees/barcode/", {"voucher_number": str(v.number), "paid_on": "2026-10-14", "how": "cash"})
    assert "already paid" in again.content.decode()
    day = client.get("/fees/collections/?from=2026-10-14").content.decode()
    assert fees["user"].get_full_name() in day  # who received it
    assert client.get(f"/fees/receipt/{payment.id}/attachment/").status_code == 200


def test_payment_attachment_is_private(fees, login_as, make_user, org, tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path
    child = fees["student"]("class-4")
    client = login_as(fees["user"])
    client.post("/fees/generate/", GENERATE)
    v = voucher_of(child)
    proof = SimpleUploadedFile("slip.pdf", b"%PDF-1.4 fake", content_type="application/pdf")
    client.post(
        "/fees/barcode/", {"voucher_number": str(v.number), "paid_on": "2026-10-14", "how": "cash", "attachment": proof}
    )
    payment = Payment.objects.get(voucher=v)
    other = login_as(make_user("principal", Scope.BRANCH, org["fsd_franchise"]))
    assert other.get(f"/fees/receipt/{payment.id}/attachment/").status_code in (403, 404)


def test_admission_voucher_screen(fees, login_as):
    child = fees["student"]("play-group")
    client = login_as(fees["user"])
    page = client.get(f"/fees/student/{child.id}/voucher/").content.decode()
    assert "Admission fee" in page and "6500" in page
    client.post(
        f"/fees/student/{child.id}/voucher/",
        {
            "first_month": "2026-10",
            "last_month": "2026-10",
            "due_on": "2026-10-20",
            "h_admission": "on",
            "a_admission": "6500",
            "h_registration": "on",
            "a_registration": "1500",
        },
    )
    assert voucher_of(child).payable == D(4300 + 6500 + 1500)
    assert "12,300" in client.get(f"/fees/student/{child.id}/").content.decode()


def test_discount_needs_someone_else_to_approve(fees, login_as, make_user):
    child = fees["student"]()
    client = login_as(fees["user"])
    client.post(
        f"/fees/student/{child.id}/discount/",
        {
            "head": FeeHead.objects.get(code="tuition").id,
            "kind": "amount",
            "value": "1000",
            "starts_on": "2026-10",
            "reason": "Hardship",
        },
    )
    c = Concession.objects.get()
    assert c.status == "pending"
    assert client.post("/fees/discounts/", {"id": c.id, "decision": "approve"}).status_code == 403
    principal = make_user("principal", Scope.BRANCH, fees["branch"])
    login_as(principal).post("/fees/discounts/", {"id": c.id, "decision": "approve"})
    c.refresh_from_db()
    assert c.status == "approved" and c.decided_by == principal


def test_who_can_see_fees(fees, org, login_as, make_user):
    child = fees["student"]()
    generate(fees)
    v = voucher_of(child)
    assert login_as(make_user("class_teacher", Scope.BRANCH, fees["branch"])).get("/fees/").status_code == 403
    other = login_as(make_user("principal", Scope.BRANCH, org["fsd_franchise"]))
    assert other.get(f"/fees/voucher/{v.id}/receive/").status_code == 404
    assert other.get(f"/fees/student/{child.id}/").status_code == 404
    assert login_as(fees["user"]).get("/fees/royalty/").status_code == 403
    ho = login_as(make_user("head_office_admin", Scope.HEAD_OFFICE, org["fsd"]))
    page = ho.get("/fees/royalty/?month=2026-10").content.decode()
    assert "Test FSD Franchise" in page and "Test LHR Franchise" not in page


def test_reports_render(fees, login_as):
    fees["student"]()
    TODAY["date"] = dt.date(2026, 9, 1)
    generate(fees, first=SEP)
    TODAY["date"] = dt.date(2026, 10, 14)
    client = login_as(fees["user"])
    assert "Child" in client.get("/fees/defaulters/").content.decode()
    assert client.get("/fees/collections/?from=2026-10-01&to=2026-10-31").status_code == 200
    assert client.get("/fees/discounts/").status_code == 200


def test_fee_pages_in_urdu(fees, login_as):
    child = fees["student"]()
    generate(fees)
    v = voucher_of(child)
    fees["user"].preferred_language = "ur"
    fees["user"].save()
    client = login_as(fees["user"])
    client.cookies["django_language"] = "ur"
    for url in ["/fees/", f"/fees/student/{child.id}/", f"/fees/voucher/{v.id}/receive/", "/fees/defaulters/"]:
        page = client.get(url, HTTP_ACCEPT_LANGUAGE="ur").content.decode()
        assert 'lang="ur"' in page and "فیس" in page, url


# --- JazzCash files --------------------------------------------------------------------------


def _xlsx(rows):
    import io

    from openpyxl import Workbook

    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    buffer.name = "jazzcash-paid.xlsx"
    return buffer


def test_voucher_list_for_jazzcash(fees, login_as):
    import io

    from openpyxl import load_workbook

    a, b = fees["student"](), fees["student"]()
    generate(fees)
    S.record_payment(b, D(4300), TODAY["date"], "counter", fees["user"], voucher=voucher_of(b))
    response = login_as(fees["user"]).get("/fees/jazzcash/vouchers.xlsx?month=2026-10")
    rows = list(load_workbook(io.BytesIO(response.content)).active.values)
    assert rows[0][0] == "Voucher ID" and len(rows) == 2  # only the unpaid voucher
    assert rows[1][0] == voucher_of(a).number and rows[1][8] == 4300


def test_jazzcash_paid_file_is_checked_then_recorded(fees, login_as):
    a, b = fees["student"](), fees["student"]()
    generate(fees)
    va, vb = voucher_of(a), voucher_of(b)
    client = login_as(fees["user"])
    upload = _xlsx(
        [
            ["JazzCash settlement report"],  # a title line above the header
            ["Sr", "Voucher No", "Amount Paid", "Transaction Date", "TID"],
            [1, va.number, 4300, "12/10/2026", "JC-1"],
            [2, 999, 1000, "12/10/2026", "JC-2"],
            [3, vb.number, "4,300", "14/10/2026", "JC-3"],
            [4, vb.number, "4,300", "14/10/2026", "JC-3"],
        ]
    )
    page = client.post("/fees/payments/import/", {"step": "check", "method": "jazzcash", "file": upload})
    body = page.content.decode()
    assert "No voucher with this ID" in body and "repeated in this file" in body
    assert not Payment.objects.exists()  # nothing saved by the check
    token = page.context["token"]
    client.post("/fees/payments/import/", {"step": "import", "token": token, "method": "jazzcash"})
    assert Payment.objects.filter(method="jazzcash").count() == 2
    va.refresh_from_db()
    assert va.status == Voucher.Status.PAID and Payment.objects.get(reference="JC-1").paid_on == dt.date(2026, 10, 12)

    upload.seek(0)  # the same file again: everything already recorded
    again = client.post("/fees/payments/import/", {"step": "check", "method": "jazzcash", "file": upload})
    assert "Already recorded" in again.content.decode() and not again.context["token"]
