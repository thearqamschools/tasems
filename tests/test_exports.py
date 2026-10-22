"""M5: every list downloads as Excel with the page's filters, and prints in full."""

import datetime as dt
import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.accounting.models import Expense, ExpenseHead
from apps.accounts.models import Scope
from apps.fees import services as S
from apps.hr.models import Employee
from apps.people.models import Enrolment, Student

pytestmark = pytest.mark.django_db


def sheet(response):
    assert response.status_code == 200
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(io.BytesIO(response.content)).active
    return [[c.value for c in row] for row in ws.iter_rows()]


@pytest.fixture
def school(org, make_user, section):
    branch = org["samanabad"]
    kids = []
    for n in range(3):
        s = Student.objects.create(
            branch=branch,
            admission_no=str(5001 + n),
            first_name=f"Pupil{n}",
            father_name=f"Father{n}",
            gender="male",
            admission_date=dt.date(2026, 4, 1),
        )
        Enrolment.objects.create(branch=branch, student=s, session=org["session"], section=section)
        kids.append(s)
    principal = make_user("principal", Scope.BRANCH, branch)
    return {"branch": branch, "kids": kids, "principal": principal, "section": section}


def test_students_excel_follows_the_filters_and_prints_every_row(school, login_as, org):
    other = Student.objects.create(
        branch=org["fsd_franchise"],
        admission_no="9",
        first_name="Elsewhere",
        gender="male",
        admission_date=dt.date(2026, 4, 1),
    )
    client = login_as(school["principal"])
    rows = sheet(client.get("/students/?format=xlsx&q=Pupil1"))
    assert rows[0][0].startswith("The Arqam Schools — ") and rows[2][0] == "Admission no."
    assert [r[1] for r in rows[3:]] == ["Pupil1"]  # only what the filter shows
    everyone = sheet(client.get("/students/?format=xlsx"))
    names = {r[1] for r in everyone[3:]}
    assert {"Pupil0", "Pupil1", "Pupil2"} <= names and other.full_name not in names  # own branch only
    assert everyone[3][3] == "Class 5" and everyone[3][4] == "A"
    page = client.get("/students/").content.decode()
    assert "format=xlsx" in page and "data-print" in page
    printed = client.get("/students/?print=1").content.decode()
    assert "data-autoprint" in printed and "Pupil2" in printed


def test_fee_lists_download(school, login_as):
    branch, user = school["branch"], school["principal"]
    S.generate_vouchers(
        branch, school["section"].session, dt.date(2026, 10, 1), dt.date(2026, 10, 1), dt.date(2026, 10, 10), user
    )
    kid = school["kids"][0]
    voucher = kid.vouchers.get()
    S.record_payment(kid, voucher.payable, dt.date(2026, 10, 5), "counter", user, voucher=voucher)
    client = login_as(user)
    vouchers = sheet(client.get("/fees/vouchers/?month=2026-10&format=xlsx"))
    assert len(vouchers) == 3 + 3 and vouchers[2][0] == "Voucher ID"
    collected = sheet(client.get("/fees/collections/?from=2026-10-05&format=xlsx"))
    assert collected[3][4] == kid.full_name and collected[3][-2] == float(voucher.payable)
    owing = sheet(client.get("/fees/defaulters/?format=xlsx"))
    assert {r[1] for r in owing[3:]} == {"Pupil1", "Pupil2"}
    families = sheet(client.get("/fees/defaulters/?by=family&format=xlsx"))
    assert families[2][0] == "Father" and len(families) == 3 + 2


def test_staff_absentees_and_expenses_download(school, login_as):
    branch, user = school["branch"], school["principal"]
    Employee.objects.create(
        branch=branch, code="E1", full_name="Teacher One", gender="female", designation="Senior Teacher",
        staff_type="teaching", joined_on=dt.date(2025, 1, 1),
    )  # fmt: skip
    Expense.objects.create(
        branch=branch, date=dt.date.today(), head=ExpenseHead.objects.get(code="utilities"), amount=Decimal(1500),
        description="Electricity bill", recorded_by=user,
    )  # fmt: skip
    client = login_as(user)
    staff = sheet(client.get("/staff/?format=xlsx"))
    assert staff[3][1] == "Teacher One" and staff[3][3] == "Senior Teacher"
    spent = sheet(client.get("/expenses/?format=xlsx"))
    assert spent[3][2] == "Electricity bill" and spent[3][7] == 1500
    absent = sheet(client.get("/attendance/absent/?format=xlsx"))
    assert absent[2][0] == "Status" and len(absent) == 3  # nobody absent: headings only


def test_excel_needs_the_same_permission_as_the_page(school, login_as, make_user, org):
    stranger = login_as(make_user("principal", Scope.BRANCH, org["lhr_franchise"]))
    response = stranger.get("/students/?format=xlsx")
    assert response.status_code in (200, 403)
    if response.status_code == 200:  # their own (empty) branch, never Samanabad's students
        assert all(r[1] not in {"Pupil0", "Pupil1", "Pupil2"} for r in sheet(response)[3:])
    teacher = login_as(make_user("class_teacher", Scope.BRANCH, school["branch"]))
    assert teacher.get("/fees/collections/?format=xlsx").status_code == 403
