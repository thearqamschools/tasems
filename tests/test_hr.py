"""Staff records, attendance from the biometric machine, and leave."""

import datetime as dt
import io

import pytest
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from apps.accounts.models import Scope
from apps.hr import services as S
from apps.hr.models import Employee, Holiday, HRSettings, LeaveType, Punch, StaffDay

pytestmark = pytest.mark.django_db
TODAY = dt.date(2026, 10, 16)  # Friday
SUN, MON, TUE, WED, THU = (dt.date(2026, 10, d) for d in (11, 12, 13, 14, 15))


@pytest.fixture(autouse=True)
def fixed_today(monkeypatch):
    moment = timezone.make_aware(dt.datetime(2026, 10, 16, 18, 0))
    monkeypatch.setattr(timezone, "localdate", lambda *a, **k: TODAY)
    monkeypatch.setattr(timezone, "now", lambda: moment)


@pytest.fixture
def hr(org, make_user):
    branch = org["samanabad"]
    HRSettings.objects.filter(branch=branch).update(grace_minutes=10, half_day_after=dt.time(10, 0))
    principal = make_user("principal", Scope.BRANCH, branch)
    counter = {"n": 0}

    def employee(**extra):
        counter["n"] += 1
        data = {
            "branch": branch,
            "code": f"E{counter['n']}",
            "machine_id": str(counter["n"]),
            "full_name": f"Staff {counter['n']}",
            "gender": "female",
            "designation": "Teacher",
            "staff_type": "teaching",
            "joined_on": dt.date(2025, 1, 1),
        }
        data.update(extra)
        return Employee.objects.create(**data)

    return {"branch": branch, "principal": principal, "employee": employee, "session": org["session"]}


def punch(e, date, hh, mm, user):
    at = timezone.make_aware(dt.datetime.combine(date, dt.time(hh, mm)))
    return Punch.objects.create(branch=e.branch, employee=e, at=at, created_by=user)


def status(e, date):
    return StaffDay.objects.get(employee=e, date=date).status


def test_status_from_punches(hr):
    on_time, late, half, absent = (hr["employee"]() for _ in range(4))
    p = hr["principal"]
    punch(on_time, WED, 7, 50, p)  # within 10 minutes' grace
    punch(on_time, WED, 14, 5, p)
    punch(late, WED, 8, 1, p)
    punch(half, WED, 10, 30, p)
    S.rebuild_days(hr["branch"], WED, WED)
    assert [status(e, WED) for e in (on_time, late, half, absent)] == ["present", "late", "half_day", "absent"]
    day = StaffDay.objects.get(employee=on_time, date=WED)
    assert (day.first_in, day.last_out, day.hours) == (dt.time(7, 50), dt.time(14, 5), 6.2)


def test_off_days_holidays_and_today(hr):
    e = hr["employee"]()
    Holiday.objects.create(date=TUE, name="Test holiday")
    S.rebuild_days(hr["branch"], SUN, TODAY)
    assert status(e, SUN) == "off" and status(e, TUE) == "holiday" and status(e, MON) == "absent"
    assert not StaffDay.objects.filter(employee=e, date=TODAY).exists()  # not in yet today
    punch(e, SUN, 9, 0, hr["principal"])
    S.rebuild_days(hr["branch"], SUN, SUN)
    assert status(e, SUN) == "present"


def test_without_duty_time_everyone_who_came_is_present(hr):
    HRSettings.objects.filter(branch=hr["branch"]).update(duty_starts=None, half_day_after=None)
    e = hr["employee"]()
    punch(e, WED, 11, 0, hr["principal"])
    S.rebuild_days(hr["branch"], WED, WED)
    assert status(e, WED) == "present"


def test_approved_leave_and_balance(hr, make_user):
    e = hr["employee"]()
    clerk = make_user("principal", Scope.BRANCH, hr["branch"])
    casual = LeaveType.objects.get(code="casual")
    casual.days_per_year = 10
    casual.save()
    req = S.request_leave(e, casual, MON, THU, "Family", clerk)
    assert req.days == 4
    with pytest.raises(ValueError):
        S.request_leave(e, casual, WED, WED, "Again", clerk)  # overlaps
    with pytest.raises(ValueError):
        S.decide_leave(req, clerk, approve=True)  # cannot approve own request
    S.decide_leave(req, hr["principal"], approve=True)
    assert status(e, TUE) == "leave"
    balance = {t.code: (allowed, used) for t, allowed, used, _p, _left in S.leave_balance(e, hr["session"])}
    assert balance["casual"] == (10, 4)
    S.cancel_leave(req, hr["principal"])
    assert status(e, TUE) == "absent"


def test_correction_is_kept_and_audited(hr):
    e = hr["employee"]()
    S.rebuild_days(hr["branch"], WED, WED)
    S.correct_day(e, WED, "present", "Machine was off", hr["principal"], first_in=dt.time(7, 40))
    S.rebuild_days(hr["branch"], WED, WED)
    assert status(e, WED) == "present"
    S.remove_correction(e, WED, hr["principal"])
    assert status(e, WED) == "absent"


DAT = b"  101\t2026-10-14 07:52:10\t1\t0\t1\t0\n  102\t2026-10-14 08:15:00\t1\t0\t1\t0\n  999\t2026-10-14 08:00:00\t1\n"


def test_machine_files_in_three_formats(hr):
    a, b = hr["employee"](machine_id="101"), hr["employee"](machine_id="102")
    check = S.check_punches(hr["branch"], S.read_machine_file(DAT, "attlog.dat"))
    assert len(check.matched) == 2 and check.unknown_ids == {"999": 1}
    assert S.save_punches(hr["branch"], check, hr["principal"], "attlog.dat") == 2
    assert status(a, WED) == "present" and status(b, WED) == "late"

    csv = b"AC-No.,Name,Time\n101,Staff,15/10/2026 07:30\n102,Staff,15/10/2026 07:31 AM\n"
    check = S.check_punches(hr["branch"], S.read_machine_file(csv, "export.csv"))
    S.save_punches(hr["branch"], check, hr["principal"], "export.csv")
    assert status(b, THU) == "present"

    wb = Workbook()
    wb.active.append(["Emp No", "Date", "Time"])
    wb.active.append([101, dt.date(2026, 10, 13), dt.time(7, 44)])
    buf = io.BytesIO()
    wb.save(buf)
    check = S.check_punches(hr["branch"], S.read_machine_file(buf.getvalue(), "x.xlsx"))
    assert len(check.matched) == 1
    again = S.check_punches(hr["branch"], S.read_machine_file(DAT, "attlog.dat"))
    assert again.already == 2  # importing the same file twice adds nothing


def test_samanabad_settings_are_seeded(seeded):
    hr = HRSettings.objects.get(branch__code="fsd-samanabad")
    assert (hr.duty_starts, hr.duty_ends, hr.working_days) == (dt.time(7, 45), dt.time(13, 45), [0, 1, 2, 3, 4, 5])
    assert not LeaveType.objects.filter(is_paid=True).exists()


def test_month_sheet_counts_unpaid_days(hr):
    e = hr["employee"]()
    punch(e, WED, 7, 30, hr["principal"])
    req = S.request_leave(e, LeaveType.objects.get(code="casual"), MON, TUE, "Family", e.branch and hr["principal"])
    from apps.accounts.models import User

    S.decide_leave(req, User.objects.create_superuser("hr-admin", password="x-Unused-123"), approve=True)
    S.rebuild_days(hr["branch"], dt.date(2026, 10, 1), TODAY)
    dates, rows = S.month_sheet(hr["branch"], dt.date(2026, 10, 1))
    totals = rows[0]["totals"]
    assert len(dates) == 31 and totals["present"] == 1 and totals["leave"] == 2
    # absent on every other working day up to Thursday 15th, plus 2 days of unpaid leave
    assert totals["unpaid"] == totals["absent"] + 2


# --- Screens ---------------------------------------------------------------------------------


def test_principal_adds_staff_and_imports_machine_file(hr, login_as):
    client = login_as(hr["principal"])
    assert client.get("/staff/").status_code == 200
    client.post(
        "/staff/new/",
        {
            "code": "T-1",
            "full_name": "Test Teacher",
            "gender": "female",
            "cnic": "33100-1234567-1",
            "phone": "0300-1234567",
            "designation": "Teacher",
            "staff_type": "teaching",
            "joined_on": "2025-04-01",
            "machine_id": "101",
        },
    )
    e = Employee.objects.get(code="T-1")
    assert (e.cnic, e.phone) == ("3310012345671", "03001234567")
    upload = io.BytesIO(DAT)
    upload.name = "attlog.dat"
    page = client.post("/staff/attendance/import/", {"step": "check", "file": upload})
    assert page.context["check"].unknown_ids == {"102": 1, "999": 1}
    client.post("/staff/attendance/import/", {"step": "import", "token": page.context["token"]})
    assert status(e, WED) == "present"
    day = client.get("/staff/attendance/?date=2026-10-14").content.decode()
    assert "Test Teacher" in day and "07:52" in day
    client.post(
        "/staff/attendance/?date=2026-10-13", {"employee": e.id, "status": "present", "note": "Machine was off"}
    )
    assert StaffDay.objects.get(employee=e, date=TUE).is_corrected
    sheet = client.get("/staff/attendance/month/?month=2026-10&format=xlsx")
    assert sheet.status_code == 200 and sheet["Content-Type"].startswith("application/vnd.openxml")


def test_staff_import_from_excel(hr, login_as):
    wb = Workbook()
    wb.active.append(
        ["Employee Code *", "Full Name *", "Gender *", "Designation *", "Staff Type *", "Joining Date *", "Machine ID"]
    )
    wb.active.append(["S-1", "Qari Sahib", "Male", "Hifz Teacher", "Qari", "01/04/2020", 7])
    wb.active.append(["S-2", "Guard", "M", "Guard", "Support", dt.date(2021, 5, 1), ""])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "staff.xlsx"
    client = login_as(hr["principal"])
    page = client.post("/staff/import/", {"step": "check", "file": buf})
    assert page.context["result"].ok
    client.post("/staff/import/", {"step": "import", "token": page.context["token"]})
    assert Employee.objects.get(code="S-1").staff_type == "hifz" and Employee.objects.get(code="S-1").machine_id == "7"


def test_self_service_leave(hr, login_as, make_user):
    teacher = make_user("class_teacher", Scope.BRANCH, hr["branch"])
    e = hr["employee"](user=teacher)
    client = login_as(teacher)
    assert "My attendance" in client.get("/staff/me/").content.decode()
    client.post(
        "/staff/me/leave/",
        {
            "leave_type": LeaveType.objects.get(code="sick").id,
            "starts_on": "2026-10-19",
            "ends_on": "2026-10-20",
            "reason": "Fever",
        },
    )
    req = e.leave_requests.get()
    assert req.days == 2 and req.status == "pending"
    assert client.get("/staff/").status_code == 403  # teachers do not see the staff register
    login_as(hr["principal"]).post("/staff/leave/", {"id": req.id, "decision": "approve"})
    req.refresh_from_db()
    assert req.status == "approved"


def test_other_branch_cannot_see_staff(hr, org, login_as, make_user):
    e = hr["employee"]()
    other = login_as(make_user("principal", Scope.BRANCH, org["fsd_franchise"]))
    assert other.get(f"/staff/{e.id}/").status_code == 404


def test_staff_import_can_skip_rows_with_problems(hr, login_as):
    wb = Workbook()
    wb.active.append(["Employee Code *", "Full Name *", "Gender *", "Designation *", "Staff Type *", "Joining Date *"])
    wb.active.append(["S-1", "Qari Sahib", "Male", "Hifz Teacher", "Qari", "01/04/2020"])
    wb.active.append(["S-2", "Guard", "Robot", "Guard", "Support", "01/05/2021"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "staff.xlsx"
    client = login_as(hr["principal"])
    page = client.post("/staff/import/", {"step": "check", "file": buf})
    assert page.context["result"].can_skip_problems
    assert "Skip problem rows and import 1 staff member" in page.content.decode()
    token = page.context["token"]
    problems = load_workbook(io.BytesIO(client.get(f"/staff/import/?problems={token}").content)).active
    assert problems.cell(2, 1).value == "S-2"
    client.post("/staff/import/", {"step": "import", "token": token, "skip_problems": "1"})
    assert Employee.objects.filter(code="S-1").exists() and not Employee.objects.filter(code="S-2").exists()
