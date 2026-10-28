"""Student register screens and Excel import."""

import datetime as dt
import io

import pytest
from openpyxl import Workbook, load_workbook

from apps.academics.models import ClassLevel, Section
from apps.accounts.models import Scope
from apps.people.importer import COLUMNS
from apps.people.models import Guardian, Student

pytestmark = pytest.mark.django_db

HEADERS = [h for _k, h, _r in COLUMNS]


def student_post(**overrides):
    data = {
        "admission_no": "A-100",
        "first_name": "Ali",
        "last_name": "Raza",
        "full_name_ur": "علی رضا",
        "gender": "male",
        "date_of_birth": "2018-05-01",
        "admission_date": "2026-04-01",
        "b_form_no": "",
        "guardian_name": "Raza Khan",
        "guardian_phone": "0300-1112223",
        "relationship": "father",
    }
    data.update(overrides)
    return data


def test_teacher_can_view_but_not_add(org, students, make_user, login_as):
    client = login_as(make_user("class_teacher", Scope.BRANCH, org["samanabad"]))
    assert client.get("/students/").status_code == 200
    assert client.get("/students/new/").status_code == 403
    assert client.get("/students/import/").status_code == 403


def test_list_shows_only_current_branch(org, students, make_user, login_as):
    client = login_as(make_user("principal", Scope.BRANCH, org["samanabad"]))
    body = client.get("/students/").content.decode()
    assert "S-1" in body and "F-1" not in body and "L-1" not in body


def test_other_branch_student_is_not_found(org, students, make_user, login_as):
    client = login_as(make_user("principal", Scope.BRANCH, org["samanabad"]))
    assert client.get(f"/students/{students['lhr_franchise'].pk}/").status_code == 404


def test_principal_adds_student_and_siblings_share_contact(org, section, make_user, login_as):
    client = login_as(make_user("principal", Scope.BRANCH, org["samanabad"]))
    response = client.post("/students/new/", student_post(section=section.pk, roll_no=1))
    assert response.status_code == 302
    ali = Student.objects.get(admission_no="A-100")
    assert ali.enrolments.get().section == section
    client.post("/students/new/", student_post(admission_no="A-101", first_name="Sara", gender="female"))
    assert Guardian.objects.filter(branch=org["samanabad"], phone="03001112223").count() == 1
    assert Guardian.objects.get(phone="03001112223").students.count() == 2


def test_form_rejects_bad_phone_and_duplicate_admission(org, students, make_user, login_as):
    client = login_as(make_user("principal", Scope.BRANCH, org["samanabad"]))
    body = client.post("/students/new/", student_post(admission_no="S-1", guardian_phone="12345")).content.decode()
    assert "already used in this branch" in body
    assert "Pakistani mobile number" in body


def test_form_rejects_girl_in_boys_section(org, make_user, login_as):
    boys = Section.objects.create(
        branch=org["samanabad"],
        session=org["session"],
        class_level=ClassLevel.objects.get(code="class-6"),
        name="A",
        gender_group="boys",
    )
    client = login_as(make_user("principal", Scope.BRANCH, org["samanabad"]))
    body = client.post("/students/new/", student_post(gender="female", section=boys.pk)).content.decode()
    assert "for Boys only" in body
    assert not Student.objects.filter(admission_no="A-100").exists()


def test_edit_student_records_history(org, students, make_user, login_as):
    client = login_as(make_user("principal", Scope.BRANCH, org["samanabad"]))
    s = students["samanabad"]
    client.post(f"/students/{s.pk}/edit/", student_post(admission_no="S-1", first_name="Bilal", status="active"))
    s.refresh_from_db()
    assert s.first_name == "Bilal"
    assert "first_name" in client.get(f"/students/{s.pk}/").content.decode()


def xlsx(rows):
    wb = Workbook()
    ws = wb.active
    ws.append([h + " *" for h in HEADERS])
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    buffer.name = "students.xlsx"
    return buffer


def row(adm, name, gender, klass, section, phone="03001234567", date=dt.date(2026, 4, 1)):
    return [adm, name, "", "", gender, "", date, "", klass, section, "", "Parent", "", phone, "Father"]


def test_template_downloads(org, make_user, login_as):
    client = login_as(make_user("principal", Scope.BRANCH, org["samanabad"]))
    response = client.get("/students/import/template.xlsx")
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.active.cell(1, 1).value == "Admission No *"


def test_import_reports_every_problem_and_saves_nothing(org, students, make_user, login_as):
    client = login_as(make_user("principal", Scope.BRANCH, org["samanabad"]))
    file = xlsx(
        [
            row("S-1", "Dup", "Male", "Class 1", "A"),  # admission no. already exists
            row("N-1", "X", "Unknown", "Class 1", "A"),  # bad gender
            row("N-2", "Y", "Male", "Class 99", "A"),  # unknown class
            row("N-3", "Z", "Female", "Class 2", "A", phone="123"),  # bad phone
            row("N-3", "Z2", "Female", "Class 2", "A"),  # repeated in file
            row("N-4", "Girl", "Female", "Class 6", "A"),
            row("N-5", "Boy", "Male", "Class 6", "A"),  # boys and girls in one Class 6 section
        ]
    )
    body = client.post("/students/import/", {"step": "check", "file": file}).content.decode()
    for text in [
        "Already exists in this branch",
        "Use Male or Female",
        "Unknown class",
        "Not a Pakistani mobile number",
        "Repeated in this file",
        "cannot share section",
    ]:
        assert text in body, text
    assert Student.objects.filter(branch=org["samanabad"]).count() == 1


def test_import_saves_all_rows_and_creates_sections(org, make_user, login_as):
    client = login_as(make_user("principal", Scope.BRANCH, org["samanabad"]))
    file = xlsx(
        [
            row("N-1", "Ahmed", "Boy", "PG", "A"),
            row("N-2", "Ayesha", "Girl", "1", "A", phone="+92 301 2345678"),
            row("N-3", "Hamza", "M", "Pre-9", "Rose", date="15/04/2026"),
            row("N-4", "Fatima", "F", "Class 6", "B"),
        ]
    )
    body = client.post("/students/import/", {"step": "check", "file": file}).content.decode()
    assert "4 students are ready to import" in body
    token = client.session["import_token"]["token"]
    response = client.post("/students/import/", {"step": "import", "token": token})
    assert response.status_code == 302
    imported = Student.objects.filter(branch=org["samanabad"])
    assert imported.count() == 4
    assert imported.get(admission_no="N-3").admission_date == dt.date(2026, 4, 15)
    assert Section.objects.get(class_level__code="pre-9", name="Rose").gender_group == "boys"
    assert Section.objects.get(class_level__code="class-6", name="B").gender_group == "girls"
    assert Section.objects.get(class_level__code="play-group", name="A").gender_group == "mixed"
    assert Guardian.objects.filter(phone="03012345678").exists()
    # The same token cannot be used twice.
    again = client.post("/students/import/", {"step": "import", "token": token})
    assert again.status_code == 302 and Student.objects.filter(branch=org["samanabad"]).count() == 4


def test_head_office_home_lists_its_branches(org, students, make_user, login_as):
    client = login_as(make_user("head_office_admin", Scope.HEAD_OFFICE, org["fsd"]))
    body = client.get("/").content.decode()
    assert "Test FSD Franchise" in body and "Test LHR Franchise" not in body


def test_branch_switcher_refuses_other_branches(org, make_user, login_as):
    client = login_as(make_user("principal", Scope.BRANCH, org["samanabad"]))
    client.post("/branch/switch/", {"branch": org["lhr_franchise"].pk})
    assert client.session.get("current_branch_id") is None


def test_health_check_needs_no_login(client):
    response = client.get("/healthz")
    assert response.status_code == 200 and response.json() == {"status": "ok"}


def test_head_office_sees_todays_registers_per_branch(org, students, make_user, login_as):
    body = login_as(make_user("head_office_admin", Scope.HEAD_OFFICE, org["fsd"])).get("/").content.decode()
    assert "Registers taken today" in body and "Diaries published today" in body


def test_father_cnic_typed_with_dashes(org, section, make_user, login_as):
    from apps.people.models import Student

    client = login_as(make_user("principal", Scope.BRANCH, org["samanabad"]))
    client.post(
        "/students/new/",
        student_post(admission_no="C-1", father_name="Abdul Rauf", father_cnic="33100-1234567-1", section=section.id),
    )
    assert Student.objects.get(admission_no="C-1").father_cnic == "3310012345671"
    body = client.post("/students/new/", student_post(admission_no="C-2", father_cnic="1234")).content.decode()
    assert "13 digits" in body


OLD_REPORT = """
<html><head><style>@page {{ size: A4 landscape; margin: 10mm; }} table {{ border-collapse: collapse; }}
td, th {{ border: 1px solid #000; padding: 3px; font-size: 9pt; }}</style></head><body>
<p>THE ARQAM SCHOOLS</p><p>Student Data</p>
<p>Count ( 3 ) Campus (Boys) Class : (Pre Nine) Section : (Rose)</p>
<table><tr><th>Comp. Code</th><th>Student Name</th><th>Father Name</th><th>Class</th><th>Section</th>
<th>Mon. Fee</th><th>Mobile</th><th>Other Number</th><th>CNIC/Occupation</th></tr>
{rows}</table></body></html>
"""


def old_report_pdf():
    from weasyprint import HTML

    rows = [
        ("9001", "Test Boy One", "Test Father One", "Pre<br>Nine", "Rose", "2900", "+923001112223<br>03211234567", "",
         "3310011111111"),
        ("9002", "Test Boy Two", "Test Father Two", "Pre<br>Nine", "Rose", "4300", "+923002223334", "2669650",
         "3310022222222<br>School"),
        ("9003", "Test Boy Three", "Test Father Three", "Pre<br>Nine", "Rose", "4000", "+923003334445", "", ""),
    ]  # fmt: skip
    html = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>" for r in rows)
    return HTML(string=OLD_REPORT.format(rows=html)).write_pdf()


def test_import_old_software_pdf_report(org, make_user, login_as):
    from apps.fees import services as fees
    from apps.fees.models import Concession, FeeHead, FeeRate

    client = login_as(make_user("principal", Scope.BRANCH, org["samanabad"]))
    upload = io.BytesIO(old_report_pdf())
    upload.name = "Student Data.pdf"
    page = client.post("/students/import/", {"step": "check", "file": upload})
    result = page.context["result"]
    assert result.ok, result.errors
    assert result.defaulted_admission_dates == 3
    client.post("/students/import/", {"step": "import", "token": page.context["token"]})
    one = Student.objects.get(admission_no="9001")
    assert (one.gender, one.father_name, one.father_cnic) == ("male", "Test Father One", "3310011111111")
    assert one.admission_date == org["session"].starts_on
    link = one.guardian_links.get()
    assert (link.guardian.phone, link.guardian.alt_phone) == ("03001112223", "03211234567")
    two = Student.objects.get(admission_no="9002")
    assert two.guardian_links.get().guardian.occupation == "School"
    assert one.enrolments.get().section.class_level.code == "pre-9"
    assert one.enrolments.get().section.gender_group == "boys"
    # the old monthly fee carries over exactly, whatever the sibling rule would say
    assert Concession.objects.get(student=one).value == 2900
    tuition = FeeHead.objects.get(code="tuition")
    head_rate = FeeRate.objects.get(branch=org["samanabad"], head=tuition)
    gross, discount, reason = fees.price(
        tuition, head_rate, org["session"].starts_on, True, [Concession.objects.get(student=one)]
    )
    assert (gross - discount, reason) == (2900, "Agreed fee")


def test_import_can_skip_rows_with_problems(org, make_user, login_as):
    client = login_as(make_user("principal", Scope.BRANCH, org["samanabad"]))
    file = xlsx(
        [
            row("N-1", "Ahmed", "Male", "Class 1", "A"),
            row("N-2", "Bilal", "Male", "Class 1", "A", phone="123"),  # bad parent mobile
            row("N-3", "Ayesha", "Female", "Class 6", "Rose"),
            row("N-4", "Usman", "Male", "Class 7", "Lily", phone="12"),  # its new section is not created
        ]
    )
    page = client.post("/students/import/", {"step": "check", "file": file})
    body = page.content.decode()
    assert "Not a Pakistani mobile number" in body
    assert "Skip problem rows and import 2 students" in body
    assert [s[1] for s in page.context["result"].good_new_sections] == ["A", "Rose"]
    token = page.context["token"]

    problems = client.get(f"/students/import/?problems={token}")
    sheet = load_workbook(io.BytesIO(problems.content)).active
    assert [sheet.cell(r, 1).value for r in (2, 3)] == ["N-2", "N-4"]
    assert "Parent Mobile" in sheet.cell(2, len(HEADERS) + 1).value

    # Without choosing to skip, nothing is imported.
    client.post("/students/import/", {"step": "import", "token": token})
    assert not Student.objects.filter(branch=org["samanabad"]).exists()

    page = client.post(
        "/students/import/",
        {
            "step": "check",
            "file": xlsx(
                [row("N-1", "Ahmed", "Male", "Class 1", "A"), row("N-2", "Bilal", "Male", "Class 1", "A", phone="123")]
            ),
        },
    )
    response = client.post(
        "/students/import/", {"step": "import", "token": page.context["token"], "skip_problems": "1"}
    )
    assert response.status_code == 302
    assert list(Student.objects.filter(branch=org["samanabad"]).values_list("admission_no", flat=True)) == ["N-1"]
    assert not Section.objects.filter(class_level__code="class-7", name="Lily").exists()


def test_missing_columns_cannot_be_skipped(org, make_user, login_as):
    client = login_as(make_user("principal", Scope.BRANCH, org["samanabad"]))
    wb = Workbook()
    wb.active.append(["Admission No", "First Name"])
    wb.active.append(["N-1", "Ahmed"])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    buffer.name = "students.xlsx"
    page = client.post("/students/import/", {"step": "check", "file": buffer})
    assert not page.context["result"].can_skip_problems
    assert "Skip problem rows" not in page.content.decode()
